from __future__ import annotations

import os
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "/Users/jiejie/Documents/lingshu-AI/output/灵枢AI内部商业化交付SOP-V1.0.docx"

BLUE = "1351A3"
NAVY = "0B2545"
MID_BLUE = "2E74B5"
LIGHT_BLUE = "EAF2FF"
PALE_BLUE = "F4F8FE"
GREEN = "16865C"
LIGHT_GREEN = "E9F6EF"
AMBER = "B36B00"
LIGHT_AMBER = "FFF4D6"
RED = "B42318"
LIGHT_RED = "FDECEC"
PURPLE = "6B4FA1"
LIGHT_PURPLE = "F1ECFA"
GRAY = "5D6875"
LIGHT_GRAY = "F2F4F7"
LINE = "CBD5E1"
WHITE = "FFFFFF"
BLACK = "17212B"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                element.set(qn("w:" + key), str(edge_data[key]))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(
                cell,
                top={"val": "single", "sz": 4, "color": LINE},
                bottom={"val": "single", "sz": 4, "color": LINE},
                left={"val": "single", "sz": 4, "color": LINE},
                right={"val": "single", "sz": 4, "color": LINE},
            )


def set_repeat_cell_text(cell, text: str, bold=False, color=BLACK, size=9.1) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int],
              header_fill=BLUE, zebra=True, font_size=9.1, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for i, h in enumerate(headers):
        set_cell_shading(hdr.cells[i], header_fill)
        set_repeat_cell_text(hdr.cells[i], h, bold=True, color=WHITE, size=9.1)
    for r_idx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            set_repeat_cell_text(
                cell, value, bold=(first_col_bold and i == 0), color=BLACK, size=font_size
            )
            if zebra and r_idx % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def set_run_font(run, size=None, bold=None, color=None, east_asia="PingFang SC") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_num_definitions(doc: Document) -> tuple[int, int, int]:
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    start_abs = max(existing_abs or [0]) + 1
    start_num = max(existing_num or [0]) + 1

    def create(abs_id: int, num_id: int, fmt: str, text: str, font=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "hybridMultilevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        lvl.append(p_pr)
        if font:
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), font)
            r_fonts.set(qn("w:hAnsi"), font)
            r_pr.append(r_fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_ref = OxmlElement("w:abstractNumId")
        abs_ref.set(qn("w:val"), str(abs_id))
        num.append(abs_ref)
        numbering.append(num)

    create(start_abs, start_num, "bullet", "•", "Arial")
    create(start_abs + 1, start_num + 1, "decimal", "%1.")
    # A separate decimal instance lets the incident-response procedure restart
    # at 1 instead of continuing the earlier weekly-production sequence.
    create(start_abs + 2, start_num + 2, "decimal", "%1.")
    return start_num, start_num + 1, start_num + 2


def set_numbering(paragraph, num_id: int, level=0) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc: Document, text: str, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    keep_with_next(p)
    return p


def add_body(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph(style="Normal")
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True, color=NAVY)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc: Document, text: str, bullet_num_id: int):
    p = doc.add_paragraph(style="Normal")
    set_numbering(p, bullet_num_id)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(doc: Document, text: str, decimal_num_id: int):
    p = doc.add_paragraph(style="Normal")
    set_numbering(p, decimal_num_id)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_callout(doc: Document, label: str, text: str, kind="blue"):
    palette = {
        "blue": (BLUE, LIGHT_BLUE), "green": (GREEN, LIGHT_GREEN),
        "amber": (AMBER, LIGHT_AMBER), "red": (RED, LIGHT_RED),
        "purple": (PURPLE, LIGHT_PURPLE), "gray": (GRAY, LIGHT_GRAY),
    }
    dark, light = palette[kind]
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent=0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, light)
    set_cell_border(cell, left={"val": "single", "sz": 20, "color": dark})
    set_cell_margins(cell, top=120, bottom=120, start=180, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}  ")
    set_run_font(r, bold=True, color=dark, size=10)
    r = p.add_run(text)
    set_run_font(r, color=BLACK, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_chapter_banner(doc: Document, number: str, title: str, subtitle: str):
    if len(doc.paragraphs) > 1:
        doc.add_page_break()
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [1550, 7810], indent=0)
    left, right = table.rows[0].cells
    set_cell_shading(left, BLUE)
    set_cell_shading(right, NAVY)
    for cell in (left, right):
        set_cell_margins(cell, top=240, bottom=240, start=220, end=220)
        set_cell_border(cell)
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(number)
    set_run_font(r, size=24, bold=True, color=WHITE)
    p = right.paragraphs[0]
    r = p.add_run(title + "\n")
    set_run_font(r, size=18, bold=True, color=WHITE)
    r = p.add_run(subtitle)
    set_run_font(r, size=9.5, color="D6E6FF")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_stage_card(doc: Document, code: str, name: str, owner: str, clock: str,
                   entry: str, action: str, output: str, pass_rule: str):
    rows = [
        ["主责 / 时限", f"{owner}｜{clock}"],
        ["准入条件", entry],
        ["关键动作", action],
        ["强制输出物", output],
        ["放行标准（DoD）", pass_rule],
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2100, 7260])
    cell = table.cell(0, 0)
    set_cell_shading(cell, BLUE)
    set_repeat_cell_text(cell, code, bold=True, color=WHITE, size=12)
    cell = table.cell(0, 1)
    set_cell_shading(cell, LIGHT_BLUE)
    set_repeat_cell_text(cell, name, bold=True, color=NAVY, size=11)
    for label, value in rows:
        row = table.add_row()
        set_repeat_cell_text(row.cells[0], label, bold=True, color=NAVY, size=9.1)
        set_cell_shading(row.cells[0], PALE_BLUE)
        set_repeat_cell_text(row.cells[1], value, size=9.1)
    set_table_geometry(table, [2100, 7260])
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_checklist(doc: Document, title: str, items: list[tuple[str, str]], owner="项目负责人"):
    add_heading(doc, title, 2)
    rows = [["□", item, evidence] for item, evidence in items]
    rows.append(["签署", f"检查人：__________  日期：__________", f"{owner}确认：__________"])
    return add_table(doc, ["状态", "检查项", "证据 / 通过标准"], rows, [620, 4670, 4070], font_size=8.8)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("内部资料 · V1.0  |  ")
    set_run_font(r, size=8, color=GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r._r.append(fld_char1)
    r._r.append(instr)
    r._r.append(fld_char2)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, MID_BLUE, 18, 10),
        ("Heading 2", 13, MID_BLUE, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("灵枢 AI  ·  商业化交付中心")
    set_run_font(r, size=8.5, bold=True, color=BLUE)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), LINE)
    borders.append(bottom)
    p_pr.append(borders)
    add_page_field(section.footer.paragraphs[0])


def add_cover(doc: Document):
    for _ in range(1):
        doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("LINGSHU AI  /  INTERNAL DELIVERY SYSTEM")
    set_run_font(r, size=9.5, bold=True, color=BLUE)
    p.paragraph_format.space_after = Pt(12)
    p = doc.add_paragraph()
    r = p.add_run("灵枢 AI\n内部商业化交付 SOP")
    set_run_font(r, size=31, bold=True, color=NAVY)
    p.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph()
    r = p.add_run("总交付流程 · 销售转交付 · 三条业务子 SOP · 治理附件")
    set_run_font(r, size=13, color=MID_BLUE)
    p.paragraph_format.space_after = Pt(18)
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360], indent=0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, BLUE)
    set_cell_margins(cell, top=180, bottom=180, start=280, end=280)
    p = cell.paragraphs[0]
    r = p.add_run("核心目标\n")
    set_run_font(r, size=10, bold=True, color="BBD7FF")
    r = p.add_run("在明确边界、明确责任、明确时间的前提下，稳定复制客户可感知的增长动作与有效询盘机会。")
    set_run_font(r, size=15, bold=True, color=WHITE)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    meta = [
        ["版本", "V1.0"], ["发布日期", "2026-09-03"],
        ["文件属性", "内部制度 / 不面向客户"], ["总负责人", "公司负责人（当前交付 DRI）"],
        ["主业务周期", "代运营 90 个自然日"], ["适用范围", "标准 SaaS、企业定制/私有化、社媒代运营"],
    ]
    add_table(doc, ["文档信息", "内容"], meta, [2100, 7260], header_fill=NAVY, zebra=False)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("CONFIDENTIAL  ·  未经授权不得对外发送")
    set_run_font(r, size=8.5, bold=True, color=RED)


def element_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def has_page_break(element) -> bool:
    return bool(element.xpath('.//w:br[@w:type="page"]'))


def is_chapter_banner(element, code: str, title: str) -> bool:
    if element.tag != qn("w:tbl"):
        return False
    rows = [child for child in element if child.tag == qn("w:tr")]
    if len(rows) != 1:
        return False
    cells = [child for child in rows[0] if child.tag == qn("w:tc")]
    return (
        len(cells) == 2
        and element_text(cells[0]) == code
        and title in element_text(cells[1])
    )


def reorder_document_sections(doc: Document) -> None:
    """Repair the interrupted draft's construction order without rewriting content blocks."""
    body = doc._element.body
    sect_pr = body.sectPr

    elements = [child for child in body.iterchildren() if child is not sect_pr]
    start = next(i for i, el in enumerate(elements) if element_text(el).startswith("0.2 三条业务线"))
    first_chapter = next(i for i, el in enumerate(elements) if is_chapter_banner(el, "01", "总交付 SOP"))
    end = first_chapter
    while end > start and has_page_break(elements[end - 1]):
        end -= 1
    management_tail = elements[start:end]
    for el in management_tail:
        body.remove(el)

    elements = [child for child in body.iterchildren() if child is not sect_pr]
    chapter_three = next(i for i, el in enumerate(elements) if is_chapter_banner(el, "03", "社媒代运营子 SOP"))
    for offset, el in enumerate(management_tail):
        body.insert(chapter_three + offset, el)

    marker_titles = {
        "00": "使用说明与管理口径",
        "01": "总交付 SOP",
        "02": "销售转交付 SOP",
        "03": "社媒代运营子 SOP",
        "04": "标准 SaaS 子 SOP",
        "05": "企业定制 / 私有化子 SOP",
        "06": "节点检查表",
        "07": "项目看板字段字典",
        "08": "异常升级与止损机制",
        "09": "交付模板与落地计划",
    }
    elements = [child for child in body.iterchildren() if child is not sect_pr]
    markers = []
    for i, el in enumerate(elements):
        for code, title in marker_titles.items():
            if is_chapter_banner(el, code, title):
                markers.append((i, code))
                break
    if len(markers) != len(marker_titles) or {code for _, code in markers} != set(marker_titles):
        raise RuntimeError(f"Chapter markers incomplete: {markers}")

    prefix = elements[:markers[0][0]]
    chunks = {}
    for idx, (start_idx, code) in enumerate(markers):
        end_idx = markers[idx + 1][0] if idx + 1 < len(markers) else len(elements)
        chunks[code] = elements[start_idx:end_idx]

    for el in list(body.iterchildren()):
        if el is not sect_pr:
            body.remove(el)
    insert_at = 0
    for el in prefix:
        body.insert(insert_at, el)
        insert_at += 1
    for code in sorted(marker_titles):
        for el in chunks[code]:
            body.insert(insert_at, el)
            insert_at += 1

    # Chapter breaks were created before the draft sections were reordered, so
    # those break-only paragraphs otherwise travel with the preceding chunk.
    # Rebuild them against the final order so every chapter starts on a fresh
    # page and a banner cannot be stranded at the foot of the prior chapter.
    for el in list(body.iterchildren()):
        if el is not sect_pr and el.tag == qn("w:p") and has_page_break(el) and not element_text(el):
            body.remove(el)

    for el in list(body.iterchildren()):
        if any(is_chapter_banner(el, code, title) for code, title in marker_titles.items()):
            page_break = OxmlElement("w:p")
            run = OxmlElement("w:r")
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "page")
            run.append(br)
            page_break.append(run)
            body.insert(body.index(el), page_break)


def main():
    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    doc = Document()
    configure_document(doc)
    bullet_id, decimal_id, incident_decimal_id = add_num_definitions(doc)
    add_cover(doc)

    # ── 阅读导航 ─────────────────────────────────────────────────────────────
    add_chapter_banner(doc, "00", "使用说明与管理口径", "先看边界与阶段门，再按业务线执行；合同与本制度冲突时立即升级。")
    add_callout(
        doc, "一页结论",
        "本 SOP 不是“做事步骤清单”，而是项目控制系统。每一阶段必须满足准入条件、形成证据、由唯一责任人签字后才能进入下一阶段。",
        "blue",
    )
    add_heading(doc, "0.1 文件目的", 1)
    add_body(doc, "建立一条可复制、可追踪、可验收、可止损的商业化交付底线，解决销售承诺与实际能力错配、账号与第三方条件拖延、生产质量不稳定、客户反馈超时以及效果口径失真等问题。")
    for item in [
        "适用于新签、续约、增购、试点、渠道转单以及重大范围变更。",
        "所有外部承诺先经过可交付性评估；所有阶段都有唯一 DRI、期限和完成定义（Definition of Done）。",
        "交付完成与商业效果分开管理：前者可承诺，后者以过程指标、机会指标和客户配合条件共同衡量。",
        "本制度严于产品手册时，以本制度作为内部执行底线；对客户以已签合同、订单及书面确认版本为准。",
    ]:
        add_bullet(doc, item, bullet_id)

    # ── 社媒代运营 ───────────────────────────────────────────────────────────
    add_chapter_banner(doc, "03", "社媒代运营子 SOP（主线）", "90 天形成可持续生产基线：80% 稳定内容 + 20% 创新实验，并建立询盘机会闭环。")
    add_heading(doc, "3.1 商业套餐、履约方式与周期", 1)
    add_callout(doc, "内部规划周期", "季框项目按合同起止日期执行；内部默认用 90 个自然日规划。前 30 日校准定位与生产模型，31–60 日建立稳定节奏并收敛有效模板，61–90 日放大有效组合并形成续约依据。该节奏属于内部目标，除非写入合同，不构成额外对外 SLA。", "blue")
    add_table(doc, ["商业套餐", "客户挂牌价", "渠道结算价", "标准产量口径"], [
        ["多平台内容矩阵｜双平台", "季框 ¥19,800", "季 ¥13,860", "每月 12–16 组母内容，按所选平台适配"],
        ["多平台内容矩阵｜三平台", "季框 ¥24,800", "季 ¥17,360", "每月 12–16 组母内容，按所选平台适配"],
        ["多平台内容矩阵｜四平台", "季框 ¥29,800", "季 ¥20,860", "每月 12–16 组母内容；预计 48–64 次发布"],
        ["单平台深度运营", "¥12,000 / 平台 / 季", "¥8,400 / 平台 / 季", "每个平台每月 12–16 条平台原生内容"],
        ["四平台深度运营", "季框 ¥48,000", "季 ¥33,600", "四个平台分别每月 12–16 条平台原生内容"],
    ], [2400, 1800, 1800, 3360], font_size=8.35, first_col_bold=True)
    add_callout(doc, "季框赠送", "签订季框赠送 1 年基础版社媒 SaaS；权益对应单个品牌，不随平台数量重复计算。价格或套餐更新时，以经批准的最新版报价和合同为准。", "green")
    add_table(doc, ["履约方式", "灵枢负责", "客户负责", "适用与时间口径"], [
        ["账号运营（标准履约）", "内容策划制作、在合规授权账号内发布、工作日基础互动、留存数据与复盘", "账号权属/验证/费用/敏感操作批准、审核、询盘销售承接", "账号在 D0 就绪时，内部目标为第 10 个工作日前首发"],
        ["内容交付（书面降级）", "选题、脚本、基础图文/视频、适配、内审、交付与复盘", "发布、互动、询盘承接和数据回传", "仅在账号未就绪、异常或客户书面选择时使用；不改变商业套餐产量"],
    ], [1900, 2810, 2520, 2130], font_size=8.25, first_col_bold=True)
    add_callout(doc, "不变的账号红线", "不提供从零注册、实名、认证、申诉、解封、养号，不接收或长期保存客户账号明文密码，不使用员工私人身份材料。任何绕过双重验证、平台规则或客户授权的要求立即停线。", "red")

    add_heading(doc, "3.2 D0 启动条件", 1)
    add_table(doc, ["全部必须满足", "内容交付模式", "账号运营模式附加项"], [
        ["商业", "合同/订单有效、启动款满足、套餐/平台/数量清晰", "同左"],
        ["资料", "品牌与产品事实、目标市场、可用素材或可执行替代素材方案", "同左"],
        ["决策", "日常接口人、审批人、询盘承接人到位；同意 2 个工作日反馈 SLA", "账号敏感操作批准人和备份联系人到位"],
        ["账号", "不要求客户账号就绪", "权属清晰、可登录、2FA 可配合、测试发布成功、无已知封禁/争议"],
        ["边界", "客户书面确认由其发布并回传数据", "客户书面授权操作范围；理解账号异常将停表/降级"],
    ], [1850, 3780, 3730], font_size=8.65, first_col_bold=True)

    add_heading(doc, "3.3 90 天里程碑", 1)
    add_table(doc, ["里程碑", "时间", "核心动作", "必须输出 / 放行标准"], [
        ["O0｜待启动", "D0 前", "完成销售交接、资料和账号审计、模式选择", "《D0 确认单》；缺口有责任人与日期"],
        ["O1｜诊断", "D0–D2", "梳理产品卖点、受众、竞品、历史数据、合规与询盘路径", "《品牌事实卡》《账号/渠道诊断》《基线数据表》"],
        ["O2｜策略冻结", "D3–D5", "确定内容支柱、平台比例、80/20、CTA、关键词、样式与周节奏", "《30 日内容策略》《首月选题池》客户确认"],
        ["O3｜首批生产", "D6–D7", "矩阵完成首 3–4 组母内容；深度运营完成各平台首批原生样本；建立 7 天库存", "内容 QA ≥85，无否决项；库存可见"],
        ["O4｜首交/首发", "最迟第 10 个工作日", "内容交付，或账号就绪时完成首发并留证", "《首批交付包》或平台发布链接/截图"],
        ["O5｜周循环", "每周", "洞察→选题→生产→QA→审批→发布/交付→互动→归因→复盘", "《周计划》《周报》《内容台账》《询盘台账》"],
        ["O6｜首月评审", "D30", "核验产量、质量、数据基线、模板表现、客户配合和风险", "矩阵 ≥12 组；深度运营每平台 ≥12 条；形成保留/优化/停用清单"],
        ["O7｜稳定放量", "D31–D60", "提升有效模板比例，连续运行 80/20，优化 CTA 与询盘承接", "稳定周节奏；未来 7 天库存；第二月复盘"],
        ["O8｜放大与验证", "D61–D83", "对有效主题、结构、场景和平台组合放大；验证新实验", "第三月数据完整；续约假设有证据"],
        ["O9｜结项", "D84–D90", "完成 90 日复盘、资产归档、下一周期计划与续约建议", "《90 日结项报告》《内容资产库》《下一周期实验表》"],
    ], [1540, 1450, 3770, 2600], font_size=8.25, first_col_bold=True)

    add_heading(doc, "3.4 80/20 内容生产系统", 1)
    add_table(doc, ["池", "占比", "来源与做法", "通过 / 退出规则"], [
        ["稳定收益池", "约 80%", "分析同赛道爆款结构、客户历史优质内容、平台趋势与已验证模板；做主题、开头、叙事、证据、CTA 的合规重组，不复制具体表达", "进入条件：至少完成内部 QA；有历史数据时优先选高于账号中位数的结构。连续 3 次显著低于基线则降级或停用"],
        ["创新实验池", "约 20%", "新题材、新钩子、新呈现方式、新人设/场景、新 CTA；默认每条只验证 1 个主要变量", "每个实验写明假设、对照、观察窗口和成功标准；双变量须批准并写明设计；到期保留、迭代或终止"],
    ], [1700, 1000, 4320, 2340], font_size=8.6, first_col_bold=True)
    add_callout(doc, "统计与冷启动", "80/20 按月度计划内容数统计：实验数=round(月计划×20%)，至少 2 条/组，其余为稳定池。前 30 日尚无客户数据时，80% 仅指符合市场共性结构的低风险基线内容，不代表已验证有效；月末必须用真实数据重排模板池。", "purple")

    add_heading(doc, "3.5 每周标准流水线", 1)
    weekly_steps = [
        ("洞察更新", "采集同赛道热门内容、平台趋势、客户销售问题、评论与询盘；记录来源和日期。"),
        ("选题排程", "从稳定池与实验池选题；标注平台、受众、目标、CTA、证据和负责人。"),
        ("内容生产", "矩阵套餐每月 12–16 组母内容滚动生产；深度运营按每个平台每月 12–16 条平台原生内容生产。"),
        ("平台适配", "根据 Instagram、Facebook、TikTok、YouTube 等平台调整比例、长度、标题、字幕、封面和 CTA。"),
        ("内部 QA", "事实、版权、合规、商业目标、表达、平台适配、CTA 七类检查；未达 85 分返修。"),
        ("客户审批", "客户在 2 个工作日内集中反馈；标准服务每组内容最多 2 轮合理修改，合同另有约定时从其约定。超时按合同默认机制或停表。"),
        ("发布 / 交付", "账号运营留存链接、时间和截图；内容交付留存文件包、版本和发送证据。"),
        ("互动与询盘", "工作日基础互动；发现询盘在台账登记并转给客户指定承接人，不代替销售成交跟进。"),
        ("复盘与沉淀", "每周标记保留/优化/停用；更新模板库、禁用词、FAQ 和下一周实验。"),
    ]
    for idx, (name, detail) in enumerate(weekly_steps, 1):
        add_number(doc, f"{name}｜{detail}", decimal_id)

    add_heading(doc, "3.6 数量、质量与验收口径", 1)
    add_table(doc, ["指标", "统一口径", "底线"], [
        ["矩阵母内容组", "围绕一个核心主题产出的主内容资产；同一母内容的平台适配不重复计为新的母内容", "每月 12–16 组，以合同为准"],
        ["深度运营原生内容", "每个平台独立选题、制作和复盘；不得用跨平台轻适配重复计为多个原生内容", "每个平台每月 12–16 条"],
        ["矩阵平台发布量", "四平台矩阵中，12–16 组母内容经适配可形成约 48–64 次发布；不等于 48–64 条独立原创", "台账中需关联母内容 ID"],
        ["计入完成", "完成生产、内部 QA、归档并按约提交客户；客户超时未审批时可记“已交付未发布”，前提是合同已明确", "不得把草稿、未 QA 或无证据内容计入"],
        ["内容质量", "每条/组单独评分：事实准确 30、合规版权 20、商业目标 20、表达 15、平台适配 10、CTA 5；各维度按满分/半分/零分计", "单条/组总分 ≥85 且无否决项"],
        ["修改", "因灵枢事实错误、漏项或不符合冻结策略的返修不计客户修改轮次；客户主观偏好/新增要求按轮次或变更处理", "标准服务每组最多 2 轮合理修改"],
    ], [1750, 5350, 2260], font_size=8.55, first_col_bold=True)
    add_callout(doc, "内容否决项", "虚假或无法核验的产品事实；未授权素材或明显侵权；违法违规、歧视、色情、危险、欺诈性表达；泄露隐私/机密；错误价格/认证/承诺；平台高风险禁词；CTA 指向错误或不可用。任一项命中即不得提交/发布。", "red")

    add_heading(doc, "3.7 进度与产能控制", 1)
    add_table(doc, ["检查点", "绿灯", "黄灯", "红灯 / 动作"], [
        ["当月产量", "月中 ≥ceil(月计划×50%)；第 3 周 ≥ceil(月计划×75%)；月末 100%", "预计低于月计划但可在 3 个工作日追回", "连续两周落后或月末未完成：停止新增实验，资源纠偏并升级"],
        ["内容库存", "始终有未来 7 天可交付/发布内容", "低于 5 天", "低于 3 天：启用备用模板和内容交付模式"],
        ["审批", "客户 ≤2 个工作日", "接近 2 日未回复", "超过 2 日：书面催办并外部等待停表；按约定默认机制处理"],
        ["返修", "一次通过率稳定；≤2 轮", "单批接近 2 轮", ">2 轮或策略反复：触发范围/决策人复核"],
        ["账号", "登录、发布和权限稳定", "验证或单平台异常", "权属/封禁/申诉/大面积失败：立即停止账号操作，切内容交付"],
    ], [1450, 3110, 2440, 2360], font_size=8.35, first_col_bold=True)

    add_heading(doc, "3.8 询盘机会闭环", 1)
    add_table(doc, ["状态", "定义", "灵枢动作", "客户动作"], [
        ["LEAD0｜互动信号", "点赞、收藏、关注、普通评论等弱意向", "记录趋势，不逐个视为询盘", "无强制"],
        ["LEAD1｜潜在线索", "询问产品、价格、交付、合作或索要资料", "服务时段内 2 小时登记并转交", "2 个工作小时内确认收到"],
        ["LEAD2｜有效询盘", "联系人/企业或需求相对明确，具备业务相关性", "补录来源、平台、时间、需求；不擅自报价", "1 个工作日内完成资格判断并回传"],
        ["LEAD3｜销售机会", "客户确认有进一步沟通、报价、样品或会议需求", "做内容归因与主题复盘", "由销售团队跟进和维护 CRM"],
        ["LEAD4｜结果", "成交、丢单、长期培育或无效", "用于内容策略复盘", "状态变化后 2 个工作日内回传"],
    ], [1450, 2920, 2690, 2300], font_size=8.45, first_col_bold=True)
    add_callout(doc, "效果判断底线", "代运营是否有效，至少同时看三层：交付健康度（数量/质量/节奏）、市场信号（完播/互动/主页访问/点击/私信）、经营机会（有效询盘/销售机会）。没有客户销售回传，不得把“无成交”单方面归因于内容。", "green")

    add_heading(doc, "3.9 周报、月报与结项", 1)
    add_table(doc, ["节奏", "必须回答的问题", "输出"], [
        ["周计划", "本周做什么、稳定与实验比例、依赖和库存是否安全？", "选题/排期/负责人/审批日/发布日"],
        ["周复盘（30–45 分钟）", "完成了什么、哪类内容高/低于基线、异常和下周调整是什么？", "《周报》+ 模板池更新"],
        ["30 日月报", "产量质量是否达标、客户配合如何、市场和询盘信号如何、成本是否健康？", "《月度运营报告》+ 下月策略"],
        ["90 日结项", "哪些内容可复制、哪些停止、询盘链路是否打通、下一阶段投入回报假设是什么？", "《90 日结项报告》+ 资产移交 + 续约建议"],
    ], [1900, 5140, 2320], font_size=8.75, first_col_bold=True)

    # ── 标准 SaaS ────────────────────────────────────────────────────────────
    add_chapter_banner(doc, "04", "标准 SaaS 子 SOP", "纯线上功能交付：快速开通、7 日激活验证、清晰划定第三方账号与 WhatsApp 边界。")
    add_heading(doc, "4.1 范围与不包含事项", 1)
    add_table(doc, ["版本", "标准能力", "挂牌价", "渠道结算价"], [
        ["基础版", "社媒板块或 WhatsApp 客服板块二选一", "季 ¥1,280 / 年 ¥3,980", "季 ¥890 / 年 ¥2,780"],
        ["旗舰版", "四平台社媒内容工厂 + WhatsApp 智能客服", "季 ¥1,980 / 年 ¥5,980", "季 ¥1,380 / 年 ¥4,180"],
    ], [1650, 3610, 2100, 2000], font_size=8.8, first_col_bold=True)
    add_table(doc, ["共同商业规则", "内部交付解释"], [
        ["7 天免费试用", "用于验证使用与基础质量，不承诺 7 天获得经营结果"],
        ["一份订阅对应一家公司下一个品牌", "多公司/多品牌不得共用；转增购或定制评估"],
        ["Token", "套餐额度覆盖常规素材、脚本、基础前端交互及客服使用；不支持额外加购 Token"],
        ["成员账号与第三方接入", "以最终合同和当期产品能力为准，不得从演示推断生产接入已包含"],
    ], [2550, 6810], font_size=8.8, first_col_bold=True)
    add_table(doc, ["套餐内包含", "标准 SaaS 不包含"], [
        ["灵枢账号开通、所购版本权限、对应板块的标准在线功能、帮助文档、标准培训与支持、系统缺陷处理", "第三方社媒账号注册/配置/认证/解封、代登录/代运营、未约定的自动发布、第三方审核、定制开发、固定经营结果"],
    ], [4680, 4680], font_size=8.8)
    add_callout(doc, "WhatsApp 销售表达", "WhatsApp 是当前客服板块支持渠道，但号码、Meta Business、模板消息、Webhook、BSP 与生产环境接通须按最终合同和接入条件核查。未完成核查时，只能承诺系统内知识库/回复能力与模拟测试，不得宣称已接通。", "red")

    add_heading(doc, "4.2 交付里程碑", 1)
    add_table(doc, ["阶段", "时限", "动作", "完成定义"], [
        ["SAAS0｜订单确认", "订单生效时点 T0", "核对套餐、周期、品牌数、使用人、开通邮箱和付款条件", "订单信息一致，无非标夹带"],
        ["SAAS1｜开通", "内部目标：条件满足后 4 个工作小时", "创建/激活灵枢账号、配置套餐权限、发送登录与边界说明", "客户可登录；套餐与期限正确"],
        ["SAAS2｜首次使用", "第 1–3 日", "提供快速上手、示例任务和 30–45 分钟标准培训/录屏", "客户至少完成一次生成、编辑或导出"],
        ["SAAS3｜7 日激活验证", "第 7 日", "检查登录、核心功能使用、问题、知识库/回复模拟和升级需求", "形成《7 日激活确认单》；未激活有原因和行动"],
        ["SAAS4｜订阅支持", "订阅期内", "处理产品使用咨询、系统缺陷、版本通知；识别定制/代运营机会", "工单留痕；范围外需求转正式评估"],
    ], [1600, 1880, 3630, 2250], font_size=8.55, first_col_bold=True)

    add_heading(doc, "4.3 D0、角色与验收", 1)
    add_table(doc, ["事项", "标准"], [
        ["SaaS D0", "订单/付款条件满足、开通邮箱有效、套餐和品牌主体明确的当天。"],
        ["产品运营", "开通权限、培训、激活追踪、边界解释和需求分流。"],
        ["技术支持", "确认灵枢系统缺陷、给出绕行方案、按异常等级修复/升级。"],
        ["客户", "保护登录信息、使用合法素材、参加培训、按指南操作并提供可复现证据。"],
        ["验收", "客户可登录；套餐权限正确；核心在线功能可完成一次闭环；边界说明已送达；严重缺陷有替代方案和时限。"],
    ], [1900, 7460], font_size=9.0, first_col_bold=True)

    add_heading(doc, "4.4 支持与卡点处理", 1)
    add_table(doc, ["卡点", "初步处理", "升级/替代"], [
        ["客户未登录/未激活", "第 2 日提醒，第 5 日再次提醒并提供短视频/图文", "第 7 日记录“客户未激活”，不视为系统交付失败；销售协同"],
        ["套餐或权限错误", "4 个工作小时内核对订单与后台", "当日修正；影响使用 >1 工作日按 P1"],
        ["功能不会用", "提供对应教程或标准培训", "超出标准培训或需代操作，转付费服务"],
        ["系统缺陷", "收集账号、时间、步骤、截图/录屏、期望和实际结果", "按 P0–P2 处理；提供可行绕行方案"],
        ["要求配置社媒账号", "重申 SaaS 边界", "可推荐客户自行完成；若需人工运营转代运营准入"],
        ["要求接通 WhatsApp", "说明只能使用系统内知识库/回复工作台与模拟测试", "转企业定制可行性评估，不保证 Meta/BSP 审批"],
    ], [1900, 3990, 3470], font_size=8.7, first_col_bold=True)

    # ── 企业定制 ─────────────────────────────────────────────────────────────
    add_chapter_banner(doc, "05", "企业定制 / 私有化子 SOP", "先冻结范围和验收，再实施；标准 30–45 个工作日，复杂项目 45–70 个工作日。")
    add_heading(doc, "5.1 适用范围与周期", 1)
    add_body(doc, "适用于菜单功能定制、工作流与权限、系统集成、私有部署、数据迁移、特定模型/云资源以及经评估的 WhatsApp 或其他第三方接口。")
    add_table(doc, ["商业基线", "V1.0 统一口径"], [
        ["最低选择", "客户至少选择 5 项定制功能"],
        ["年度服务价", "max（所选菜单挂牌价合计，¥19,800）；菜单合计不超过 ¥19,800 时仍按 ¥19,800"],
        ["系统权益", "已包含旗舰版系统使用权益，不另行收费"],
        ["渠道结算", "成交价 70% 后向下取整至十位；第三方资源按实际采购成本"],
        ["菜单版本", "报价必须附当期批准的 26 项菜单及版本；复杂逻辑、非标接口或专项部署另行评估定价"],
        ["服务期", "包含 90 天深度陪跑及合同期内技术运维；二者范围按合同与本 SOP 分开管理"],
    ], [2400, 6960], font_size=8.75, first_col_bold=True)
    add_heading(doc, "5.1.1 定制菜单价格快照（2026 年 8 月版）", 2)
    add_table(doc, ["类别", "功能菜单", "建议挂牌价", "渠道结算价"], [
        ["AI/内容", "客户专属提示词与品牌知识配置", "¥1,800–2,800", "¥1,260–1,960"],
        ["AI/内容", "AI 智能体定制", "¥3,800–4,800", "¥2,660–3,360"],
        ["AI/内容", "内容生产工作流", "¥3,800–4,800", "¥2,660–3,360"],
        ["AI/内容", "审核、通知或任务流", "¥1,800–2,800", "¥1,260–1,960"],
        ["AI/内容", "多语言生成脚本", "¥1,800–2,800", "¥1,260–1,960"],
        ["视频/数字人", "Gemini 系模型成片生产工作流", "¥3,800–4,800", "¥2,660–3,360"],
        ["视频/数字人", "Seedance 系模型成片生产工作流", "¥3,800–4,800", "¥2,660–3,360"],
        ["视频/数字人", "长文本、多模态等多模型控制集群", "¥1,800–2,800", "¥1,260–1,960"],
        ["视频/数字人", "客户专属数字人形象（本地部署）", "¥3,800–4,800", "¥2,660–3,360"],
        ["视频/数字人", "数字人视频生产工作流（本地部署）", "¥4,800 起", "¥3,360 起"],
        ["视频/数字人", "数字人直播工作流", "¥4,800 起", "¥3,360 起"],
        ["治理/分析", "多品牌或多区域工作空间", "¥1,800–3,800", "¥1,260–2,660"],
        ["治理/分析", "自定义角色与审批权限", "¥1,800–3,800", "¥1,260–2,660"],
        ["治理/分析", "SSO 单点登录", "¥1,800–3,800", "¥1,260–2,660"],
        ["治理/分析", "社媒监听与竞品看板", "¥3,800–4,800", "¥2,660–3,360"],
        ["治理/分析", "达人及 UGC 管理模块", "¥3,800–4,800", "¥2,660–3,360"],
        ["治理/分析", "评论区维护与舆情管理看板", "¥4,800 起", "¥3,360 起"],
        ["治理/分析", "获客归因与经营看板", "¥3,800–4,800", "¥2,660–3,360"],
        ["WhatsApp", "WhatsApp 客服流程深度定制", "¥3,800–4,800", "¥2,660–3,360"],
        ["WhatsApp", "WhatsApp 线索分配与跟进", "¥2,800–3,800", "¥1,960–2,660"],
        ["集成/部署", "自定义数据看板", "¥1,800–3,800", "¥1,260–2,660"],
        ["集成/部署", "品牌白标与专属域名", "¥1,800–3,800", "¥1,260–2,660"],
        ["集成/部署", "数据迁移与初始化", "¥1,800–3,800", "¥1,260–2,660"],
        ["集成/部署", "客户专属云实例", "¥4,800 实施费+云资源", "¥3,360+云资源成本"],
        ["集成/部署", "企业级 SLA 和容灾配置", "¥2,800–4,800", "¥1,960–3,360"],
        ["集成/部署", "复杂 CRM、ERP 或电商集成", "¥3,800–4,800/个", "¥2,660–3,360/个"],
    ], [1550, 3970, 1940, 1900], font_size=7.7, first_col_bold=False)
    add_callout(doc, "菜单控制", "上表用于销售交接核对，不替代正式报价。任何指导价范围、终身使用、本地部署、云资源或接口数量必须在合同中写清；发布新版产品手册时同步更新本快照版本。", "amber")
    add_table(doc, ["类型", "建议交付周期", "适用条件"], [
        ["标准定制", "内部目标 30–45 个工作日", "从 D0 至上线；边界明确、环境稳定、接口少、客户决策快"],
        ["复杂集成 / 私有化", "内部目标 45–70 个工作日", "从 D0 至上线；多系统、多环境、数据迁移、安全审查或复杂权限"],
        ["未知能力 / 强第三方依赖", "先做 5–15 个工作日付费调研或 PoC", "WhatsApp/平台 API 审核未确定、接口文档不全、技术可行性未证实"],
        ["上线后深度陪跑", "90 个自然日", "已交付范围内的使用辅导、调优、缺陷闭环和效果分析"],
        ["合同期技术运维", "陪跑结束至合同期满", "已交付系统的缺陷处理、可用性维护和合同约定支持；新增需求走变更"],
    ], [1900, 2300, 5160], font_size=8.75, first_col_bold=True)
    add_callout(doc, "工期与 PoC 口径", "上述工期为内部规划目标，不含 G0/C0、另行 PoC 和外部等待；C2–C7 可在风险可控时并行，以批准后的综合计划为准。PoC 必须先冻结假设、样本、预算上限、成功阈值和退出条件，输出结论并明确转正式、继续验证或终止。除非写入合同，不作为额外对外 SLA。", "blue")

    add_heading(doc, "5.2 里程碑与阶段门", 1)
    add_table(doc, ["里程碑", "建议时限", "关键动作", "输出 / 放行标准"], [
        ["C0｜售前评估", "1–3 个工作日", "业务/技术/安全/成本/第三方初判，确定 S/M/L/X", "《可行性评估》《初步范围》《风险与假设》"],
        ["C1｜启动", "D0–D2", "建立团队、环境和计划；确认数据、接口和决策人", "《项目章程》《D0 确认》《环境清单》"],
        ["C2｜需求澄清", "5–10 个工作日", "用户流程、角色权限、数据、接口、非功能、异常与验收场景", "《需求规格说明》无关键待定项"],
        ["C3｜范围冻结", "3–5 个工作日", "确认原型/流程/字段/验收/排除项/变更基线", "客户签署《范围基线》《UAT 计划》"],
        ["C4｜实施", "10–30 个工作日", "开发、配置、集成、部署；版本、代码、数据和成本受控", "可测试版本；单元/配置自测通过"],
        ["C5｜SIT", "5–10 个工作日", "系统集成、权限、异常、安全、性能与回归测试", "P0/P1 缺陷为 0；P2 有处置计划"],
        ["C6｜UAT", "5–10 个工作日", "客户按冻结场景验收；记录缺陷、变更和签署", "关键场景 100% 通过；客户书面验收"],
        ["C7｜上线", "3–5 个工作日", "备份、迁移、切换、验证、培训、回滚预案", "上线检查通过；责任与监控到位"],
        ["C8｜90 日陪跑", "上线后 D1–D90", "前 4 周每周、后 8 周每两周跟进；月度总结", "遗留问题闭环；资产移交；结项签署"],
    ], [1450, 1600, 3770, 2540], font_size=8.25, first_col_bold=True)

    add_heading(doc, "5.3 定制项目 D0", 1)
    for item in [
        "合同、首款与采购/安全前置条件满足；项目范围和报价基础一致。",
        "客户业务负责人、技术负责人、验收人和决策人到位；例会与反馈 SLA 确认。",
        "最低需求资料、现有流程、接口文档、数据样例、部署环境信息可用。",
        "第三方依赖已标明由谁申请、谁付费、预期审核周期和失败替代方案。",
        "私有化项目明确服务器/云账户、网络、安全策略、访问方式、备份与上线窗口。",
        "WhatsApp 等接入仅在完成账号/商务主体/模板/BSP/API 可行性核查并签署专项范围后计入 D0。",
    ]:
        add_bullet(doc, item, bullet_id)

    add_heading(doc, "5.4 需求与验收设计", 1)
    add_table(doc, ["必须冻结", "最小内容"], [
        ["业务范围", "用户角色、使用场景、流程、输入输出、异常路径、排除项"],
        ["功能范围", "菜单/模块、页面/字段、权限、配置项、通知和日志"],
        ["数据与接口", "来源、格式、频率、映射、权限、错误码、限流、留存和删除"],
        ["非功能", "性能、可用性、兼容性、安全、审计、备份、恢复和容量"],
        ["验收", "测试场景、测试数据、预期结果、通过阈值、缺陷等级、签署人和反馈窗口"],
        ["成本", "人工、第三方模型/云/API 资源、预计用量、80% 预警与超额处理"],
    ], [2100, 7260], font_size=9.0, first_col_bold=True)
    add_callout(doc, "第三方资源费", "包年方案含模型 API、Token 及共享云资源 ¥200/月、全年 ¥2,400；按月使用、不跨月累计。达到月额度 80% 时预警，未经书面确认不自动超额。WhatsApp 消息、付费外部数据、客户专属云实例及专项第三方服务另行结算。具体合同有差异时按签署版本执行。", "amber")

    add_heading(doc, "5.5 AI 客服 / 回复工作台专项 QA", 1)
    add_body(doc, "凡涉及知识库、客服回复、WhatsApp 场景或其他渠道回复能力，上线前必须完成场景集测试。标准场景不少于 50 条，复杂项目不少于 100 条。")
    add_table(doc, ["场景类别", "示例", "V1.0 内部放行标准"], [
        ["产品事实", "规格、MOQ、交付、价格、认证、售后", "关键事实错误 0；回答来源覆盖 100%"],
        ["不确定 / 无答案", "资料缺失、超范围、政策不明", "不得编造；正确拒答或转人工 100%"],
        ["异议与敏感", "竞品、投诉、退款、法律、隐私", "未经授权的高风险承诺 0"],
        ["多轮与改写", "上下文追问、不同说法、多语言", "语义一致；关键约束不丢失"],
        ["人工接管", "高意向、投诉、复杂售前、异常", "应转人工场景识别 100%"],
        ["总体可用性", "全场景人工评审", "可用回答率 ≥95%，且全部否决项为 0"],
    ], [1900, 3380, 4080], font_size=8.7, first_col_bold=True)

    add_heading(doc, "5.6 上线与 90 日陪跑边界", 1)
    add_table(doc, ["包含", "不包含 / 需变更"], [
        ["已交付范围内的培训、使用辅导、参数/提示词调优、缺陷修复、使用/效果分析、遗留问题闭环和资产移交", "新增模块/接口/品牌/语言、持续内容代运营、大规模新数据迁移、账号注册认证申诉、未约定 WhatsApp 接入、7×24 值守、无限次重做"],
    ], [4680, 4680], font_size=9.0)
    add_body(doc, "陪跑建议：第 1–4 周每周例会；第 5–12 周每两周一次；每 30 日形成使用与问题报告；第 90 日完成结项。新增需求全部进入变更流程。")

    add_heading(doc, "5.7 WhatsApp / 第三方接入专项边界", 1)
    add_table(doc, ["灵枢可承诺", "不能无条件承诺", "不可行时替代"], [
        ["按已批准接口文档实施约定连接、配置知识库/回复逻辑、完成沙箱与约定测试、提供技术记录", "Meta/BSP/平台审核一定通过、号码一定可用、模板消息一定获批、第三方政策不变、历史账号一定能恢复", "系统内模拟测试；人工导入导出；Webhook/邮件中转；先完成知识库工作台；切换到客户已就绪的合规账号"],
    ], [3300, 3200, 2860], font_size=8.65)
    add_callout(doc, "停线条件", "账号权属/数据合法性不明、要求规避平台审批、客户拒绝安全措施、关键环境不可访问、UAT 标准持续变化、第三方成本超预算或重大缺陷未关闭时，项目必须停线并升级。", "red")

    # ── 节点检查表 ───────────────────────────────────────────────────────────
    add_chapter_banner(doc, "06", "节点检查表", "每项都要有证据；“做过”不等于“通过”。可直接复制到项目主档使用。")
    add_callout(doc, "签署规则", "表内每一项按“通过 / 不适用 / 未通过”标记；不适用必须写理由。存在红线项时不得整体签署通过。电子签署、项目群书面确认或系统审批均可作为证据。", "blue")

    add_checklist(doc, "6.1 G0 商机准入检查表", [
        ("客户主体、品牌、目标市场和真实业务目标清楚", "客户沟通记录 / 商机准入单"),
        ("选择的业务线、套餐、平台/模块、数量和周期明确", "报价方案 / 范围摘要"),
        ("基线数据或“无基线”已记录；未承诺固定经营结果", "基线表 / 边界说明"),
        ("账号权属、可用性、第三方审批与 WhatsApp 诉求已识别", "账号预审 / 技术预审"),
        ("素材版权、产品事实、知识库和客户配合能力已初判", "输入清单"),
        ("S/M/L/X 分级、预计人工、第三方成本和毛利风险已评估", "分级与成本表"),
        ("关键风险有替代路径；X 项目已拒绝或转 PoC", "风险台账 / 决策记录"),
        ("拟承诺的期限经交付确认", "排期确认"),
    ], "总负责人")

    add_checklist(doc, "6.2 G1 销售转交付检查表", [
        ("合同、订单、报价和回款条件版本唯一", "链接 / 版本号"),
        ("所有口头、聊天、渠道附加承诺已逐条录入", "承诺差异表"),
        ("客户联系人、决策人、验收人、询盘承接人完整", "干系人表"),
        ("范围、排除项、数量、质量、期限、修改轮次一致", "范围摘要"),
        ("账号、材料、接口、环境缺口有责任人与截止日", "待补清单"),
        ("渠道项目的催办、解释、变更审批和直联权限明确", "渠道分工表"),
        ("项目负责人已选择接受 / 有条件接受 / 拒绝", "交接会议纪要"),
        ("D0 条件与首个里程碑已确定", "初始计划"),
    ], "销售 + 项目负责人")

    add_checklist(doc, "6.3 G2 启动与 D0 检查表", [
        ("合同/付款/授权已满足业务线最低条件", "商务确认"),
        ("交付模式与账号边界书面确认", "D0 确认单"),
        ("内部 DRI、客户接口人、决策人和验收人到位", "项目章程"),
        ("必要输入已齐，剩余缺口不阻断首阶段", "输入台账"),
        ("项目群、主档、看板、会议节奏和命名规范建立", "系统链接"),
        ("双时钟、反馈 SLA、停表和变更机制已说明", "启动会纪要"),
        ("里程碑、风险、成本和验收基线已登记", "项目计划 / 风险台账"),
        ("客户书面确认 D0 日期", "群记录 / 签署"),
    ])

    add_checklist(doc, "6.4 账号就绪检查表（账号运营适用）", [
        ("账号为客户合法拥有，客户控制注册邮箱/手机号", "客户声明 / 管理后台截图"),
        ("管理员权限、业务资产和平台映射正确", "权限截图"),
        ("2FA 正常，验证责任人与备份联系人可联系", "测试记录"),
        ("测试登录与测试发布成功", "链接 / 截图 / 时间"),
        ("无已知封禁、所有权争议、异常支付或高风险提示", "账号健康记录"),
        ("灵枢使用正式、最小必要授权，不共用 root 权限", "授权记录"),
        ("未接收/留存明文密码；敏感信息按安全方式交换", "项目负责人确认"),
        ("客户知悉注册、实名、认证、申诉、解封不在范围", "边界签收"),
        ("异常时切内容交付或暂停的路径已确认", "备选方案"),
    ], "项目负责人 + 客户")

    add_checklist(doc, "6.5 G3 策略 / 范围冻结检查表", [
        ("目标、受众/用户、场景和主要成功指标清楚", "策略/需求说明"),
        ("范围内与范围外事项逐条列明", "范围基线"),
        ("内容支柱/功能流程/字段权限等核心方案完成", "策略稿 / 原型"),
        ("事实来源、禁用词、版权、数据与合规要求明确", "事实卡 / 合规清单"),
        ("数量、节奏、质量阈值、验收方法和修改轮次明确", "验收计划"),
        ("开放问题无关键阻断；未决项有责任人与日期", "问题清单"),
        ("客户/验收人完成书面确认", "签署 / 群确认"),
        ("后续新增需求将走变更流程", "变更基线"),
    ])

    add_checklist(doc, "6.6 内容生产与 QA 检查表", [
        ("内容卡包含母内容 ID、平台、目标、受众、CTA、稳定/实验标签", "内容台账"),
        ("产品、价格、认证、交付等事实有来源", "事实引用"),
        ("素材权利和授权可追溯", "素材授权 / 来源"),
        ("无违法违规、隐私泄露、误导承诺和平台禁用风险", "合规检查"),
        ("开头、结构、证据和 CTA 与本次目标一致", "QA 评分"),
        ("尺寸、时长、字幕、标题、封面、链接符合平台要求", "平台适配检查"),
        ("QA 总分 ≥85 且无否决项", "QA 单"),
        ("版本、文件、交付/发布证据完整", "归档链接"),
    ], "内容负责人 + QA")

    add_checklist(doc, "6.7 AI 客服 / 回复能力 QA 检查表", [
        ("知识库来源、版本、责任人与更新时间明确", "知识库清单"),
        ("标准项目 ≥50 条、复杂项目 ≥100 条测试场景", "场景集"),
        ("产品、MOQ、价格、认证、交付、售后等关键事实无错误", "测试结果"),
        ("无答案时不编造并能正确拒答/转人工", "无答案场景"),
        ("投诉、法律、隐私、退款等高风险内容无越权承诺", "敏感场景"),
        ("应转人工场景识别率 100%", "转人工测试"),
        ("来源覆盖 100%，总体可用回答率 ≥95%", "评审记录"),
        ("WhatsApp 仅按已签范围标记为模拟或实际接入", "范围与环境证据"),
    ], "QA + 产品/技术")

    add_checklist(doc, "6.8 私有部署 / 上线检查表", [
        ("部署架构、环境、域名/证书、网络和访问控制确认", "部署方案"),
        ("数据分类、迁移、备份、恢复和删除策略确认", "数据计划"),
        ("SIT/UAT 结果满足标准，P0/P1 缺陷为 0", "测试报告"),
        ("第三方密钥/账号使用客户或正式项目凭据", "凭据登记（不含明文）"),
        ("容量、监控、日志、告警和 80% 成本预警可用", "运维检查"),
        ("上线窗口、负责人、业务验证与回滚方案批准", "上线计划"),
        ("培训、管理员和支持渠道到位", "培训签到 / 指南"),
        ("上线后观察期和 90 日陪跑节奏明确", "陪跑计划"),
    ], "技术负责人 + 客户")

    add_checklist(doc, "6.9 G6 客户验收检查表", [
        ("提交版本为内部 QA 通过的唯一版本", "版本号 / QA 单"),
        ("交付清单、使用说明、测试/发布证据完整", "交付包"),
        ("客户验收人与反馈截止时间明确", "发送记录"),
        ("反馈按事实错误、范围内缺陷、主观修改、新增需求分类", "客户反馈单"),
        ("范围内问题已修复并复验", "缺陷关闭记录"),
        ("范围外反馈已转变更并书面确认", "变更单"),
        ("遗留项有责任人、期限和影响说明", "遗留项表"),
        ("客户书面验收或满足合同视为验收条件", "验收单 / 证据"),
    ])

    add_checklist(doc, "6.10 月度 / 结项检查表", [
        ("本周期里程碑、数量、质量和使用指标完整", "月报 / 结项报告"),
        ("双时钟、延误原因、客户配合和第三方等待如实呈现", "进度记录"),
        ("经营信号与交付结果分层；统计口径一致", "数据字典"),
        ("异常、变更、返工和成本偏差已复盘", "台账 / 复盘"),
        ("有效模板/知识/配置/代码和禁用项已沉淀", "资产库"),
        ("客户未决事项、权限、账号和数据完成移交/回收", "移交清单"),
        ("续约、增购、优化或终止建议有证据", "建议书"),
        ("总负责人完成结项审阅", "结项签署"),
    ], "项目负责人 + 总负责人")

    # ── 项目看板字段 ─────────────────────────────────────────────────────────
    add_chapter_banner(doc, "07", "项目看板字段字典", "让进度、依赖、质量、成本和经营信号在一张表上被看见。")
    add_heading(doc, "7.1 看板最小结构", 1)
    add_table(doc, ["视图", "用途", "默认筛选 / 排序"], [
        ["项目总览", "查看每个客户当前阶段、灯号、下个里程碑和 DRI", "在服项目；按灯号→下个截止日"],
        ["待办与阻塞", "查看逾期、外部等待、缺输入、需决策事项", "红/黄；按升级等级→等待天数"],
        ["内容生产", "从选题到发布/交付的内容流水线", "本周；按状态→计划发布日期"],
        ["定制实施", "需求、版本、测试、缺陷、上线与陪跑", "开放项；按里程碑→缺陷等级"],
        ["异常与变更", "所有 P0–P3 事件和范围变更", "未关闭；按等级→响应截止"],
        ["经营与续约", "使用、市场、询盘、成本和续约信号", "按客户→最近月度"],
    ], [1750, 4870, 2740], font_size=8.8, first_col_bold=True)
    add_callout(doc, "字段纪律", "项目主表一项目一行；内容、询盘、异常、变更、测试和成本作为关联子表。所有日期使用实际日期字段，不用自由文本。所有枚举由管理员维护，禁止随意新建同义状态。", "blue")

    board_fields = [
        ("基础 P001–P012", [
            ["P001", "项目编号", "文本/唯一", "必填", "例如 LS260901"],
            ["P002", "客户主体", "关联/文本", "必填", "合同相对方"],
            ["P003", "品牌/项目名称", "文本", "必填", "可与客户不同"],
            ["P004", "渠道伙伴", "关联", "条件", "直签为空"],
            ["P005", "业务线", "单选", "必填", "代运营/SaaS/定制私有化"],
            ["P006", "交付模式", "单选", "必填", "内容交付/账号运营/线上自助/实施"],
            ["P007", "复杂度", "单选", "必填", "S/M/L/X"],
            ["P008", "总负责人", "人员", "必填", "当前公司负责人"],
            ["P009", "项目 DRI", "人员", "必填", "唯一责任人"],
            ["P010", "销售/渠道负责人", "人员", "必填", "商务接口"],
            ["P011", "客户日常接口人", "联系人", "必填", "含联系方式"],
            ["P012", "客户决策/验收人", "联系人", "必填", "可为两人"],
        ]),
        ("范围 S001–S012", [
            ["S001", "合同/订单链接", "URL", "必填", "最终版本"],
            ["S002", "套餐/服务包", "单选/文本", "必填", "与订单一致"],
            ["S003", "合同开始/结束日", "日期", "必填", "商业周期"],
            ["S004", "品牌数", "数字", "必填", "一个订阅原则上一品牌"],
            ["S005", "平台/模块", "多选", "必填", "社媒或系统功能"],
            ["S006", "国家/语言", "多选", "必填", "防止隐性扩围"],
            ["S007", "计划数量", "数字+单位", "必填", "母内容/发布/功能点"],
            ["S008", "修改轮次", "数字", "必填", "原则上 ≤2"],
            ["S009", "范围内摘要", "长文本", "必填", "一句话可复述"],
            ["S010", "范围外摘要", "长文本", "必填", "客户已知晓"],
            ["S011", "验收方法", "文本/链接", "必填", "可测量"],
            ["S012", "特殊承诺", "长文本", "条件", "含口头/渠道承诺"],
        ]),
        ("时间 T001–T014", [
            ["T001", "项目状态", "单选", "必填", "待准入/待启动/进行中/等待/暂停/结项"],
            ["T002", "当前阶段门", "单选", "必填", "G0–G7 或业务里程碑"],
            ["T003", "D0 日期", "日期", "条件", "条件齐备当日"],
            ["T004", "内部计划开始/结束", "日期", "必填", "可控时钟"],
            ["T005", "客户期望日期", "日期", "条件", "不等于承诺"],
            ["T006", "合同交付日期", "日期", "条件", "如有"],
            ["T007", "下个里程碑", "文本/单选", "必填", "可执行"],
            ["T008", "下个截止日", "日期时间", "必填", "提醒依据"],
            ["T009", "里程碑完成率", "百分比", "必填", "按 DoD"],
            ["T010", "外部等待开始", "日期时间", "条件", "停表起点"],
            ["T011", "外部等待原因", "单选+文本", "条件", "客户/平台/第三方"],
            ["T012", "累计外部等待天数", "公式", "自动", "与内部工期分开"],
            ["T013", "状态灯", "公式/单选", "必填", "绿/黄/红"],
            ["T014", "最后更新/下次更新", "日期时间", "必填", "防止看板失真"],
        ]),
        ("输入与账号 A001–A012", [
            ["A001", "输入完整度", "百分比", "必填", "按业务线清单"],
            ["A002", "缺失输入", "多选/文本", "条件", "责任人与日期另列"],
            ["A003", "客户反馈 SLA", "数字", "必填", "默认 2 个工作日"],
            ["A004", "最近催办/下次催办", "日期", "条件", "外部等待必填"],
            ["A005", "账号权属", "单选", "条件", "清晰/待证/争议"],
            ["A006", "账号就绪状态", "单选", "条件", "未提供/核查中/就绪/异常"],
            ["A007", "管理员与授权", "单选", "条件", "最小必要权限"],
            ["A008", "2FA/验证联系人", "联系人", "条件", "含备份"],
            ["A009", "测试登录/发布", "单选+证据", "条件", "通过/失败/不适用"],
            ["A010", "账号风险", "多选", "条件", "封禁/认证/支付/权属等"],
            ["A011", "WhatsApp 状态", "单选", "条件", "不适用/模拟/评估/沙箱/生产"],
            ["A012", "账号异常替代模式", "单选", "条件", "内容交付/暂停/换账号/专项评估"],
        ]),
        ("内容 C001–C016", [
            ["C001", "母内容 ID", "文本/唯一", "必填", "关联平台适配"],
            ["C002", "内容主题/支柱", "单选+文本", "必填", "策略一致"],
            ["C003", "稳定/实验", "单选", "必填", "80/20"],
            ["C004", "实验假设", "文本", "条件", "实验内容必填"],
            ["C005", "目标平台/受众", "多选", "必填", "平台适配依据"],
            ["C006", "内容目标", "单选", "必填", "认知/互动/询盘等"],
            ["C007", "CTA", "文本", "必填", "路径可用"],
            ["C008", "来源/参考链接", "URL", "条件", "洞察可追溯"],
            ["C009", "素材授权状态", "单选", "必填", "已核验/待补/不可用"],
            ["C010", "生产负责人", "人员", "必填", "唯一"],
            ["C011", "内容状态", "单选", "必填", "选题/制作/内审/客户审/待发/已发/归档"],
            ["C012", "QA 分数/否决项", "数字+多选", "必填", "≥85 且否决项 0"],
            ["C013", "计划/实际交付发布日", "日期", "必填", "双日期"],
            ["C014", "客户反馈轮次", "数字", "条件", "新增需求分开"],
            ["C015", "发布/交付证据", "URL", "必填", "链接或文件包"],
            ["C016", "数据观察与结论", "数值+文本", "条件", "保留/优化/停用"],
        ]),
        ("质量与实施 Q/D001–Q/D012", [
            ["Q001", "知识库版本/来源", "文本/URL", "条件", "客服项目"],
            ["Q002", "测试场景数", "数字", "条件", "标准≥50/复杂≥100"],
            ["Q003", "关键事实错误数", "数字", "条件", "目标 0"],
            ["Q004", "可用回答率", "百分比", "条件", "目标 ≥95%"],
            ["Q005", "转人工正确率", "百分比", "条件", "目标 100%"],
            ["Q006", "来源覆盖率", "百分比", "条件", "目标 100%"],
            ["D001", "需求/范围版本", "文本/URL", "条件", "定制项目"],
            ["D002", "环境状态", "单选", "条件", "开发/测试/生产"],
            ["D003", "接口/第三方状态", "单选", "条件", "文档/沙箱/审核/生产"],
            ["D004", "当前版本/构建", "文本", "条件", "可追溯"],
            ["D005", "P0/P1/P2 缺陷数", "数字", "条件", "放行依据"],
            ["D006", "SIT/UAT/上线状态", "单选", "条件", "关联报告"],
        ]),
        ("风险、成本与效果 R/F/E001–R/F/E015", [
            ["R001", "最高风险等级", "单选", "必填", "P0–P3/无"],
            ["R002", "关键风险/阻塞", "长文本", "条件", "一句话描述"],
            ["R003", "风险责任人与截止日", "人员+日期", "条件", "不可空"],
            ["R004", "升级状态", "单选", "条件", "未升级/已升级/决策中/关闭"],
            ["R005", "替代/恢复方案", "文本", "条件", "红黄项必填"],
            ["F001", "合同金额/回款状态", "金额+单选", "必填", "商业健康"],
            ["F002", "预算人工/实际人工", "小时", "必填", "效率与毛利"],
            ["F003", "第三方预算/实际成本", "金额", "条件", "达到 80% 预警"],
            ["F004", "变更金额/工期影响", "金额+天", "条件", "关联变更单"],
            ["E001", "交付完成率", "百分比", "必填", "确定性结果"],
            ["E002", "质量一次通过率", "百分比", "必填", "过程质量"],
            ["E003", "使用/发布健康度", "百分比", "条件", "业务线定义"],
            ["E004", "有效询盘数", "数字", "条件", "按统一定义"],
            ["E005", "销售机会数/回传率", "数字/百分比", "条件", "客户责任"],
            ["E006", "续约健康度", "单选", "月度", "绿/黄/红+理由"],
        ]),
    ]
    for title, rows in board_fields:
        add_heading(doc, title, 2)
        add_table(doc, ["字段 ID", "字段名称", "类型", "要求", "说明 / 枚举"], rows, [1100, 2350, 1400, 1000, 3510], font_size=8.15, first_col_bold=True)

    add_heading(doc, "7.2 核心公式与自动提醒", 1)
    add_table(doc, ["规则", "建议公式 / 触发", "自动动作"], [
        ["里程碑逾期", "当前时间 > 下个截止日 且状态非完成", "黄灯；超过 2 个工作日转红并通知总负责人"],
        ["外部等待", "存在等待开始且未结束", "冻结内部承诺倒计时；每日累计等待天数"],
        ["反馈超时", "提交客户后 >2 个工作日无反馈", "催办、标记外部等待；按合同默认机制或暂停"],
        ["内容库存", "库存覆盖天数=已通过 QA 的可发库存÷未来 14 日计划日均发布量", "不足 5 天黄灯，不足 3 天红灯"],
        ["月度进度", "已合格母内容组 / 月计划", "月中 <50% 或第 3 周 <75% 预警"],
        ["成本预警", "实际第三方成本 / 预算 ≥80%", "暂停非必要调用并发起费用确认"],
        ["长期未更新", "最后更新 >2 个工作日", "提醒项目 DRI；>3 日通知总负责人"],
    ], [1800, 4700, 2860], font_size=8.7, first_col_bold=True)

    add_heading(doc, "0.2 三条业务线与默认边界", 1)
    add_table(doc, ["业务线", "默认交付形态", "建议周期", "内部红线"], [
        ["社媒代运营（主线）", "商业套餐=内容矩阵/单平台深度运营；标准履约=账号运营；内容交付仅为书面降级", "按季框合同起止；内部用 90 日规划；账号就绪时内部目标第 10 个工作日前首发", "不代注册、认证、申诉、实名；账号未就绪不得承诺发布"],
        ["标准 SaaS", "基础版二选一；旗舰版含四平台内容工厂与 WhatsApp 智能客服", "内部目标 4 个工作小时开通 + 7 日激活验证；订阅期支持", "第三方实际接入以合同与核查为准；不夹带代运营"],
        ["企业定制 / 私有化", "至少 5 项；年费=max（菜单合计，¥19,800），含旗舰版权益", "内部目标：标准 30–45、复杂 45–70 个工作日；上线后 90 日陪跑", "第三方 API、WhatsApp、私有环境先验证可行性；外部等待单列"],
    ], [1950, 3030, 2050, 2330], font_size=8.65, first_col_bold=True)
    add_callout(doc, "账号与 WhatsApp 统一口径", "标准 SaaS 只交付线上功能，可表述为“AI 客服知识库与回复工作台 / WhatsApp 场景回复生成与模拟测试”。不得表述为已接通 WhatsApp。代运营仅操作客户自有、权属清晰、可正常登录且已正式授权的账号；定制项目中的实际接入须单独立项。", "red")

    add_heading(doc, "0.3 颜色与阅读方法", 1)
    add_table(doc, ["标识", "含义", "要求"], [
        ["蓝色｜阶段门", "里程碑、标准动作、必备输出", "按顺序执行，未通过不得跳转"],
        ["绿色｜通过 / 可复制", "已验收、指标健康、可进入下一阶段", "沉淀证据和模板"],
        ["黄色｜预警 / 决策", "尚未停线，但需当日或限时纠偏", "明确责任人、截止时间、替代方案"],
        ["红色｜红线 / 停线", "重大风险、越权承诺、合规或质量否决项", "立即停止、升级、保全证据"],
        ["紫色｜实验", "20% 创新内容、技术验证或未规模化能力", "单独标注，不污染稳定交付口径"],
        ["灰色｜说明", "定义、背景或参考信息", "不作为单独放行依据"],
    ], [1800, 3600, 3960], font_size=8.7, first_col_bold=True)

    add_heading(doc, "0.4 文档控制", 1)
    add_table(doc, ["事项", "规则"], [
        ["生效与变更", "由总负责人批准。涉及价格、合同、结果承诺、账号边界、P0/P1 响应标准的变更须升级审批。"],
        ["项目级裁剪", "可删减不适用动作，但不得删除阶段门、验收证据、风险升级、账号红线和变更控制。"],
        ["命名规范", "客户简称_业务线_项目编号_文件类型_日期_版本，例如：ABC_代运营_LS260901_月报_20261003_V1.0。"],
        ["证据保存", "合同/订单、会议纪要、确认记录、交付包、验收结果、异常单和变更单进入项目主档；聊天口头承诺须 1 个工作日内书面化。"],
        ["复审频率", "每季度复审一次；发生 P0、连续两次 P1、重大产品变化或新业务定价时即时复审。"],
    ], [2200, 7160], font_size=9.0, first_col_bold=True)

    add_heading(doc, "0.5 快速目录", 1)
    add_table(doc, ["章节", "内容", "谁优先阅读"], [
        ["01", "总交付 SOP：G0–G7、进度与治理", "所有参与人"],
        ["02", "销售转交付 SOP：准入、交接包、承诺冲突", "销售 / 渠道 / 项目负责人"],
        ["03", "社媒代运营子 SOP：90 天主线、80/20 内容系统、询盘闭环", "内容 / 运营 / QA"],
        ["04", "标准 SaaS 子 SOP：即时开通、7 日激活、支持边界", "销售 / 客服 / 产品运营"],
        ["05", "企业定制与私有化子 SOP：范围冻结、实施、UAT、90 日陪跑", "方案 / 产品 / 技术"],
        ["06–09", "节点检查表、看板字段、异常升级表、模板清单", "项目负责人 / 交付治理"],
    ], [1050, 5170, 3140], font_size=8.8, first_col_bold=True)

    # ── 总交付 SOP ───────────────────────────────────────────────────────────
    add_chapter_banner(doc, "01", "总交付 SOP", "一条主流程、八个阶段门、双时钟与统一停线规则。")
    add_heading(doc, "1.1 管理原则", 1)
    principles = [
        ("先准入，后承诺", "未完成复杂度、账号、合规、资源和第三方依赖评估，不得锁定交期或效果。"),
        ("先定义完成，再开始动作", "每个里程碑必须预先写清输入、动作、输出物、验收人和放行标准。"),
        ("一个结果，一个 DRI", "协作人可以多名，但每个阶段只有一名最终负责者。"),
        ("双时钟", "同时展示“内部可控工期”和“外部等待时间”；不可隐去客户或第三方造成的停表。"),
        ("范围—时间—成本联动", "任何新增需求都必须在这三者中至少调整一项，并留下变更记录。"),
        ("质量优先于形式进度", "未通过否决项，即使到期也不得伪装为完成。"),
        ("交付结果与经营效果分层", "按约交付内容/功能属于确定结果；播放、粉丝、询盘、成交属于共同影响结果，不做无条件保证。"),
        ("异常有替代路径", "账号不就绪时切内容交付；接口不可用时切模拟/人工；反馈超时按默认机制推进或停表。"),
    ]
    add_table(doc, ["原则", "执行解释"], [[a, b] for a, b in principles], [2350, 7010], font_size=9.0, first_col_bold=True)

    add_heading(doc, "1.2 关键角色与 RACI", 1)
    add_body(doc, "符号：A=最终负责；R=执行；C=会签/咨询；I=知会。每项只能有一个 A。S/M 项目按授权矩阵下放；L/X、红线、超预算和补偿事项上收总负责人。授权人缺席时必须有书面代理，不得多人共同兜底。")
    raci_rows = [
        ["可交付性准入", "A", "C", "R", "C", "C", "C", "I"],
        ["销售交接 / D0", "I", "A", "R", "I", "C", "I", "C"],
        ["策略 / 范围冻结", "I", "A", "C", "R", "C", "C", "C"],
        ["生产 / 实施", "I", "A", "I", "R", "C", "R", "C"],
        ["内部质量验收", "I", "R", "I", "C", "A", "C", "I"],
        ["客户验收 / 上线", "I", "A", "C", "R", "C", "R", "R"],
        ["范围变更", "C", "A", "C", "C", "C", "C", "C"],
        ["红色异常决策", "A", "R", "I", "C", "C", "R", "I"],
        ["月度复盘 / 续约建议", "I", "A", "C", "R", "C", "C", "C"],
    ]
    add_table(doc, ["事项", "总负责人", "项目负责人", "销售/渠道", "内容/实施", "QA", "产品/技术", "客户"], raci_rows, [2180, 980, 1120, 1120, 1050, 800, 1130, 980], font_size=7.9, first_col_bold=True)

    add_heading(doc, "1.3 项目复杂度分级", 1)
    add_table(doc, ["等级", "典型特征", "处理方式"], [
        ["S｜标准", "单品牌、标准套餐、材料齐、无第三方复杂接入、审批链短", "按标准 SLA 执行；允许模板化批量交付"],
        ["M｜增强", "多平台/双语/轻度定制/素材需补齐/单一外部依赖", "增加澄清与专项 QA；总负责人知会"],
        ["L｜复杂", "多品牌多国家、私有化、多接口、复杂权限、强合规或紧交期", "先方案评审再报价；单独排期、预算与风险金"],
        ["X｜实验", "能力未验证、强依赖平台审批、希望保证经营结果或明显超出现有产品", "不得进入标准合同；只可做有退出条件的付费 PoC 或拒绝"],
    ], [1400, 5150, 2810], font_size=8.8, first_col_bold=True)
    add_callout(doc, "分级维度", "品牌数、平台数、国家/语言、内容量、材料成熟度、知识库质量、接口与部署、审批链、期限、人工量、第三方成本、账号/合规风险。任一维度达到 L，应按 L 管理。", "gray")

    add_heading(doc, "1.4 八个阶段门（G0–G7）", 1)
    add_stage_card(doc, "G0", "商机与可交付性准入", "总负责人或授权交付负责人（A）；销售（R）", "标准 1 个工作日；复杂 2–5 个工作日", "客户目标、拟购产品、预算/期限、账号/材料/系统情况和非标承诺已初步采集。", "完成 S/M/L/X 分级；识别账号、WhatsApp、接口、合规、资源、效果承诺与成本风险；确认建议交付模式。", "《商机准入单》《边界说明》《初步风险清单》《建议报价/周期》。", "范围可描述、成本可测算、关键依赖有替代路径；否则进入条件接受或拒绝。")
    add_stage_card(doc, "G1", "销售正式转交付", "项目负责人", "签约/回款条件满足后 1 个工作日内", "合同、订单、报价、销售承诺和客户联系人完整；G0 已通过。", "召开结构化交接会；逐条复述范围、数量、期限、账号边界、非标承诺和风险。", "《销售转交付单》《承诺差异表》《待客户补充清单》《初始里程碑》。", "交付确认“接受/有条件接受”；所有口头承诺已书面化；不得带病启动。")
    add_stage_card(doc, "G2", "项目启动与 D0 生效", "项目负责人", "满足业务线 D0 条件后当日", "首款/授权/材料/账号或环境达到最低启动条件；关键联系人到位。", "建立项目群、看板和主档；召开启动会；确认沟通、审批、停表、变更与验收机制。", "《项目章程》《启动会纪要》《D0 确认单》《详细计划》《风险台账》。", "客户与内部对交付模式、时间基线、责任边界和首个里程碑无歧义。")
    add_stage_card(doc, "G3", "策略 / 方案与范围冻结", "项目负责人（A）；内容/实施负责人（R）", "代运营 2–5；定制 8–15 个工作日", "必要洞察、产品资料、品牌口径、知识库、技术环境和决策人反馈齐备。", "形成内容策略、功能清单、流程/原型、验收指标与排除项；处理开放问题。", "《策略/需求说明》《范围基线》《验收指标》《变更基线》。", "客户书面确认；未确认部分有默认处理或明确暂停；新增需求进入变更。")
    add_stage_card(doc, "G4", "生产 / 配置 / 实施", "项目负责人（A）；内容/技术负责人（R）", "按批准计划滚动执行", "G3 已通过；生产素材、接口、权限、测试数据和任务排期就绪。", "按标准作业单执行；每日更新看板；每周检查产能、质量、库存、依赖和成本。", "内容包/系统版本、过程记录、测试记录、交付清单、变更记录。", "计划完成率达阈值；无未处理否决项；输出可追溯且版本唯一。")
    add_stage_card(doc, "G5", "内部验收", "QA（A）；项目负责人（R）", "每批次/版本提交客户前", "交付物已完成自检；相关事实来源、素材版权、功能日志与测试证据可查。", "按风险矩阵全检或抽检；记录缺陷、返修和复验。", "《内部 QA 单》《缺陷单》《复验记录》《可提交版本》。", "代运营内容 ≥85 分且无否决项；客服/系统测试达到阈值；P0/P1=0，P2 仅可在不阻断且有责任人、期限和批准时遗留。")
    add_stage_card(doc, "G6", "客户验收 / 发布 / 上线", "项目负责人", "客户反馈窗口原则上 2 个工作日；以合同为准", "G5 通过；客户验收人和发布/上线条件到位。", "提交交付包；收集结构化反馈；按约定轮次修改；账号运营须留存发布证据。", "《交付确认单》《客户反馈单》《上线/发布记录》《遗留项清单》。", "客户书面验收，或满足合同约定的视为验收条件；遗留项有责任人与期限。")
    add_stage_card(doc, "G7", "稳定运营、复盘与结项", "项目负责人（A）；总负责人（I/红线时决策）", "周复盘 / 30 日月复盘 / 周期结束", "生产/系统已进入稳定期；数据口径和经营信号可采集。", "复盘质量、效率、询盘/使用、异常、成本与客户配合；提出下一周期实验、优化或续约建议。", "《周报》《月报》《90 日结项/陪跑报告》《资产移交清单》《续约建议》。", "交付物、数据、风险和资产完整归档；客户未决事项闭环；完成结项签署。")

    add_heading(doc, "1.5 D0、双时钟与状态灯", 1)
    add_table(doc, ["控制项", "定义", "看板要求"], [
        ["D0", "业务线最低启动条件全部满足的当天。签约日、建群日、付款日均不自动等于 D0。", "记录 D0 条件、满足日期、确认人；未满足则状态为“待启动”。"],
        ["Dn / 工作日", "D1、D2 等默认指从 D0 起的自然日；凡按工作日计算，必须明确写“工作日”。T0 仅用于 SaaS 订单生效时点。", "项目章程记录客户时区、工作日历和节假日；禁止混用。"],
        ["内部工期", "由灵枢可控的生产、配置、开发、测试时间。", "按里程碑承诺；内部延误不得伪装为客户等待。"],
        ["外部等待", "客户审批/材料/账号、平台审核、第三方接口或不可控环境造成的等待。", "单独累计；写明起止、催办记录、影响与替代路径。"],
        ["绿灯", "里程碑按计划；未来 7 天输入齐；无高风险。", "继续执行。"],
        ["黄灯", "预计延误 ≤2 个工作日、单项输入临期、质量/库存/成本接近阈值。", "当日提出纠偏；24 小时内明确负责人和恢复点。"],
        ["红灯", "关键路径延误 >2 个工作日、合规/安全/账号权属、承诺冲突、重大质量或预算失控。", "暂停受影响动作并按 P0/P1 升级；仅安全/合规/权属或未受控重大风险触发全项目停线。"],
    ], [1600, 4890, 2870], font_size=8.75, first_col_bold=True)

    add_heading(doc, "1.6 范围变更与结果口径", 1)
    for item in [
        "任何新增品牌、平台、语言、模块、接口、内容量、修改轮次、发布频次、实时支持或提前交期，均需提交《变更申请单》。",
        "项目负责人在 1 个工作日内评估对范围、排期、费用、资源、质量和既有里程碑的影响；客户/销售确认前不执行。",
        "紧急变更可先采取止损动作，但须在 24 小时内补齐记录和授权。",
        "确定性结果只包括合同内且可控的交付物、功能、数量、质量阈值和服务动作。播放、粉丝、询盘、成交等作为经营信号，不承诺固定数值。",
        "若确需对经营指标作阶段目标，必须写明基线、统计平台、归因窗口、排除项、客户责任和未达时的优化动作，而非无条件退款或保证。",
    ]:
        add_bullet(doc, item, bullet_id)
    add_callout(doc, "禁止跳步", "任何人不得以“客户着急”“老板口头同意”“先做再说”为由绕过 G0、G1、G3、G5 或红色停线。紧急不等于无记录。", "red")

    # ── 销售转交付 ──────────────────────────────────────────────────────────
    add_chapter_banner(doc, "02", "销售转交付 SOP", "把销售语言翻译为可交付范围；交接不完整，项目不启动。")
    add_heading(doc, "2.1 目标与时限", 1)
    add_body(doc, "销售转交付不是“把客户拉进群”，而是完成责任、范围、承诺、风险和证据的正式转移。签约/回款前完成 G0，满足启动条件后 1 个工作日内完成 G1。")
    add_callout(doc, "唯一入口", "所有新签、续约、增购、渠道转单、试点转正式项目，均必须通过同一份《销售转交付单》。未登记在交接单中的承诺，不得默认为交付义务；发现遗漏仍须立即升级处理。", "blue")

    add_heading(doc, "2.2 售前准入资料", 1)
    add_table(doc, ["类别", "必采信息", "判定重点"], [
        ["客户与目标", "主体、品牌、市场、联系人、决策人、目标、现状基线、预算、期望日期", "目标是否可定义；决策链是否真实；交期是否可行"],
        ["产品与范围", "业务线、套餐、平台/模块、数量、语言、地区、服务方式、验收口径", "是否落入标准范围；是否存在隐藏人工"],
        ["内容与资料", "品牌资料、产品事实、素材、禁用词、历史内容、竞品参考、知识库", "真实性、版权、完整度和可持续供给"],
        ["账号与授权", "账号权属、可登录性、管理员、2FA、限制状态、授权方式、发布权限", "是否满足账号运营模式；否则切内容交付"],
        ["系统与接口", "现有系统、部署环境、接口文档、沙箱、数据类型、第三方审批", "标准功能还是定制；数据/安全风险和外部依赖"],
        ["商业与资源", "价格、回款、第三方费用、预计人工、渠道分工、毛利风险", "是否可盈利；是否需要资源锁定或专项报价"],
    ], [1800, 4850, 2710], font_size=8.65, first_col_bold=True)

    add_heading(doc, "2.3 三条业务线专项准入", 1)
    add_table(doc, ["业务线", "必须额外确认", "未满足时处理"], [
        ["代运营", "内容交付/账号运营模式；平台与数量；客户 2 个工作日反馈；素材权利；账号就绪；询盘接收人", "账号未就绪切内容交付；素材不足先建 7 天库存；不承诺代注册/认证"],
        ["标准 SaaS", "套餐、品牌数、使用人、开通邮箱、培训联系人、线上功能边界", "第三方账号配置/WhatsApp 接通诉求转定制评估；不得夹带人工运营"],
        ["定制/私有化", "需求边界、部署环境、接口责任、验收测试、数据分类、第三方资源、上线窗口", "资料或环境不明则先付费调研/PoC；不得直接承诺固定上线日"],
    ], [1700, 4840, 2820], font_size=8.7, first_col_bold=True)

    add_heading(doc, "2.4 销售交接包（强制 18 项）", 1)
    handoff_items = [
        ("最终合同/订单/报价版本", "文件可打开，关键条款已标注"),
        ("客户主体与开票/回款状态", "满足合同启动条件"),
        ("客户、决策人、验收人与日常接口人", "姓名、角色、联系方式完整"),
        ("S/M/L/X 分级及理由", "总负责人已批准 M/L/X"),
        ("业务线与交付模式", "代运营须明确内容交付或账号运营"),
        ("品牌、平台、语言、地区、模块、数量", "与报价和合同一致"),
        ("所有时间承诺", "区分期望日期、合同日期、内部计划"),
        ("全部口头/聊天承诺", "逐条附证据与可交付性判断"),
        ("范围外事项与排除项", "客户已知晓"),
        ("账号/授权/2FA/限制状态", "不得保存明文密码"),
        ("素材、产品事实、知识库、数据", "缺口、责任人和截止日明确"),
        ("接口/部署/第三方条件", "文档、沙箱、审批和费用明确"),
        ("验收方法与修改轮次", "可执行、可测量"),
        ("客户配合义务", "审批 SLA、提供资料、询盘承接人"),
        ("风险、假设与替代路径", "红黄项已提示"),
        ("历史项目/投诉/敏感事项", "交付已知晓"),
        ("预计人工、外采与第三方成本", "不突破毛利底线"),
        ("渠道分工与客户沟通权限", "谁催办、谁解释、谁有变更批准权"),
    ]
    add_checklist(doc, "交接清单", handoff_items, "销售 + 项目负责人")

    add_heading(doc, "2.5 交接会议", 1)
    add_table(doc, ["项目", "标准"], [
        ["必到人员", "销售/渠道、项目负责人；M/L 项目增加内容或技术负责人、QA；必要时总负责人参加。"],
        ["建议时长", "S：30 分钟；M：45 分钟；L：60–90 分钟。"],
        ["固定议程", "客户目标与现状 → 合同范围 → 数量/质量/期限 → 账号与第三方 → 口头承诺 → 风险 → D0 条件 → 首个里程碑。"],
        ["会议输出", "接受 / 有条件接受 / 拒绝；条件项必须包含责任人、截止日和逾期处理。"],
        ["完成定义", "交接包 18 项齐全；承诺冲突已解决；项目负责人书面接收；D0 尚未满足时明确为“待启动”。"],
    ], [1850, 7510], font_size=9.0, first_col_bold=True)

    add_heading(doc, "2.6 销售禁止承诺清单", 1)
    forbidden = [
        "固定播放量、涨粉量、询盘数、成交数、爆款概率或保证 ROI。",
        "为客户从零注册、实名、认证、申诉、解封、养号，或使用灵枢人员身份资料。",
        "标准 SaaS 包含第三方账号配置、自动发布、WhatsApp 实际接通、人工代运营。",
        "代运营默认包含上门拍摄、广告投放费、达人费、复杂特效、销售跟进、节假日实时值守、重大舆情处置或无限修改。",
        "未经评估的 API、私有部署、多语言、多品牌、多国家或迁移项目的固定上线日。",
        "绕过平台规则、版权、隐私、数据安全或客户内部审批。",
    ]
    for item in forbidden:
        add_bullet(doc, item, bullet_id)
    add_callout(doc, "承诺冲突处理", "发现合同/报价/聊天与可交付能力冲突时，项目负责人当天登记红色异常；销售负责还原证据，总负责人决定修订、增补费用、调整工期、提供替代方案或拒绝。未经决策，不得让交付团队自行“兜底”。", "red")

    add_heading(doc, "2.7 渠道项目额外规则", 1)
    for item in [
        "交接单同时记录终端客户与渠道方联系人，明确谁是合同相对方、谁验收、谁付款。",
        "明确渠道一线响应、客户催办、需求解释、变更审批和客诉升级的责任边界。",
        "渠道对客户的额外承诺必须逐条列入“渠道承诺差异表”，不得以未在灵枢合同中为由忽略风险。",
        "未经书面授权，灵枢不得跨过渠道直接改变商业条款；涉及质量、安全或合规风险时可直接通知必要责任人。",
    ]:
        add_bullet(doc, item, bullet_id)

    # ── 异常升级与止损 ─────────────────────────────────────────────────────
    add_chapter_banner(doc, "08", "异常升级与止损机制", "异常先控影响、再定责任；每个等级都有响应时限、升级人和恢复条件。")
    add_heading(doc, "8.1 异常分级与响应 SLA", 1)
    add_table(doc, ["等级", "判定示例", "首次响应", "决策 / 恢复目标", "必须通知"], [
        ["P0｜重大", "数据泄露或越权、违法违规、账号权属争议、错误内容大面积发布、生产系统全面不可用、对外重大舆情", "15 分钟内止损并通知", "1 小时内形成处置组和外部口径；持续更新直至恢复", "总负责人、项目 DRI、技术/QA、客户决策人；必要时法务"],
        ["P1｜严重", "关键里程碑预计延误 >2 个工作日、账号被限/封、关键功能不可用、关键事实错误已对外、预算预计超 100%", "2 个工作小时内", "当日给出替代方案、责任人和恢复点", "总负责人、项目 DRI、相关负责人、客户接口人"],
        ["P2｜一般", "单批 QA 不通过、局部发布失败、客户反馈超时、非关键缺陷、库存低于 5 天、成本达到 80%", "1 个工作日内", "2 个工作日内纠偏或纳入批准计划", "项目 DRI、执行人；必要时客户接口人"],
        ["P3｜轻微", "文案小错、非关键数据缺失、单次流程偏差、体验优化建议", "2 个工作日内", "进入待办并在周复盘关闭", "执行人、项目 DRI"],
    ], [900, 3190, 1350, 2420, 1500], font_size=7.9, first_col_bold=True)
    add_callout(doc, "等级取高", "同一事件命中多个等级时按最高级处理。首次响应是确认、止损和指定负责人，不等于承诺已经修复。所有 P0/P1 必须建立独立异常单。", "red")
    add_callout(doc, "计时日历", "默认服务时段为双方项目章程约定时区的工作日 09:00–18:00；上述响应时限从监控发现或收到有效报告起算。非服务时段从下一个值守窗口起算，除非合同另购 7×24 值守。任何已被人员发现的数据安全、违法违规或账号劫持风险仍应立即止损并通知。", "amber")

    add_heading(doc, "8.2 标准处置闭环", 1)
    for text in [
        "发现与保全｜记录时间、项目、版本、链接/截图、影响范围和报告人；不得先删证据再汇报。",
        "立即止损｜暂停发布、自动任务、模型调用、接口、账号操作或上线切换；必要时撤回错误内容。",
        "分级与建单｜项目 DRI 在首次响应时限内确认 P0–P3、牵头人、更新频率和参与人。",
        "影响核验｜区分已发生、可能发生和未受影响范围；核查客户、平台、数据、成本与合同影响。",
        "方案决策｜选择修复、回滚、切换内容交付、人工兜底、范围变更、延期、补交或终止。",
        "内外沟通｜只发布经批准的事实、影响、临时措施和下次更新时间；不猜测、不甩锅、不越权赔偿。",
        "验证恢复｜由独立于修复人的 QA 按恢复清单复验；修复人和项目 DRI 不得单独解除 P0/P1 停线。",
        "复盘关闭｜P0 在 3 个工作日内、P1 在 5 个工作日内完成根因、行动项、责任人和制度更新；30 日后验证防复发效果。",
    ]:
        add_number(doc, text, incident_decimal_id)

    add_heading(doc, "8.3 高频卡点决策表", 1)
    add_table(doc, ["卡点", "立即动作", "恢复 / 替代路径", "升级条件"], [
        ["销售承诺超范围", "冻结相关工作，核对合同、报价、聊天和成本", "修订承诺、增补费用/工期、提供等价替代或拒绝", "当天无法达成内部决策；已对客户造成损失"],
        ["客户材料不足", "列最小缺口、责任人、截止日并启动外部等待", "使用已批准替代素材；缩小主题；先交不依赖材料部分", "阻断关键路径 >2 个工作日"],
        ["客户审批超时", "到期当日催办并记录停表", "按合同默认批准/顺延/暂停；继续准备无风险库存", "连续两批超时或影响合同日期"],
        ["账号登录/发布异常", "停止反复登录和敏感操作，保全错误信息", "切内容交付；客户自行恢复/申诉；换合规就绪账号", "权属争议、封禁、疑似安全事件"],
        ["内容 QA 未通过", "禁止提交/发布；按否决项返修", "改用稳定模板或降低产量，优先恢复 7 天库存", "连续两批 <85 分或同类否决项复发"],
        ["客服答案错误", "暂停自动回复或切人工；隔离错误知识", "修正来源、重测相关场景并全量回归关键事实", "错误已发送客户、涉及价格/认证/法律/隐私"],
        ["第三方接口/审核失败", "记录供应商状态，不承诺平台恢复时间", "沙箱/模拟、人工导入导出、Webhook/邮件中转或延期", "无替代路径且阻断验收"],
        ["效果信号偏弱", "先核验发布量、受众、CTA、归因和销售回传", "一次只改 1–2 个变量，运行有截止日的修正实验", "连续两个观察窗无改善；需增加预算或扩围"],
        ["成本接近/超过预算", "80% 预警时暂停非必要消耗，复核重试和模型档位", "降档、限额、改人工/批处理或书面增购", "预计超 100% 或未经批准已超额"],
    ], [1750, 2600, 3050, 1960], font_size=7.95, first_col_bold=True)

    add_heading(doc, "8.4 停线与恢复门槛", 1)
    add_table(doc, ["必须停线", "恢复必须同时满足"], [
        ["违法违规、侵权、隐私/数据安全、账号权属不清；客户要求绕过平台或审批；P0 未控制；关键事实/价格/认证无法核验；无可用回滚路径。P1 默认暂停受影响动作，只有影响不可隔离时全项目停线。", "根因或直接风险已隔离；责任人与批准人到位；修复/替代已完成；QA 按场景复验通过；数据/账号/预算状态清楚；客户与内部收到统一恢复口径；新的里程碑和观察期已登记。"],
    ], [4680, 4680], font_size=8.8)
    add_callout(doc, "补偿权限", "项目团队可以提出补交、延长服务、费用调整或终止建议，但不得自行对外承诺退款、赔偿、免费扩围或固定经营结果。所有补偿由总负责人与商务依据合同批准。", "amber")

    add_heading(doc, "8.5 外部等待老化与资源释放", 1)
    add_table(doc, ["累计等待", "项目状态与动作", "商业 / 排期结果"], [
        ["0–2 个工作日", "正常外部等待；记录缺口、责任人、催办日和无风险可继续项", "保留原排期"],
        ["第 3–5 个工作日", "黄灯；再次催办并启用替代路径", "提示里程碑风险，不承诺追回全部等待"],
        ["超过 5 个工作日", "红灯；暂停受阻动作并释放预留人员/算力", "恢复后重新进入容量队列，不自动保留原日期"],
        ["超过 10 个工作日", "发起正式重基线或变更", "重新确认范围、费用、交付日和客户责任"],
        ["超过 20 个工作日", "由商务和总负责人决定继续、长期暂停、阶段结项或按合同终止", "形成书面决策并完成资产/权限处理"],
    ], [1750, 4610, 3000], font_size=8.45, first_col_bold=True)

    add_heading(doc, "8.6 风险型 QA 与 CAPA", 1)
    add_table(doc, ["对象", "检查方式", "批次失败 / 放行规则"], [
        ["否决项：事实、版权、合规、隐私、价格/认证、CTA/链接", "100% 检查", "任一命中即整批暂停并全检"],
        ["新客户、新模板、20% 实验、敏感行业、重大活动", "100% 人工复核", "QA 独立放行；不得由制作人自批"],
        ["连续 3 批稳定且低风险的模板", "每批抽检 ≥30% 且不少于 3 条；自动规则仍 100%", "抽样任一否决项或两条低于 85 分，整批转全检"],
        ["客服关键事实/转人工、高风险场景", "关键场景 100%；其余按批准测试集", "关键错误 0；P0/P1=0"],
        ["定制 SIT/UAT", "关键路径与安全场景 100%；其余按测试计划", "P0/P1=0；P2 仅凭风险接受单遗留"],
    ], [2380, 3580, 3400], font_size=8.35, first_col_bold=True)
    add_callout(doc, "CAPA", "每个 P0/P1 建立纠正与预防措施：根因类别、立即纠正、系统性预防、适用项目、负责人、到期日、验证样本、30 日验证日和有效/无效结论。措施逾期或同根因复发时自动重开并升级。", "purple")

    # ── 模板目录与落地计划 ─────────────────────────────────────────────────
    add_chapter_banner(doc, "09", "交付模板与落地计划", "把 SOP 变成每天可执行的表单、看板和会议节奏；先在 3 个项目试运行。")
    add_heading(doc, "9.1 强制模板目录", 1)
    add_table(doc, ["模板", "使用阶段 / 负责人", "最小字段", "完成标志"], [
        ["T01 商机准入单", "G0 / 销售", "客户目标、业务线、范围、基线、账号/接口、预算期限、S/M/L/X、风险", "总负责人或授权交付负责人批准/拒绝"],
        ["T02 销售转交付单", "G1 / 销售+项目 DRI", "18 项交接包、承诺差异、范围排除、客户角色、D0 缺口", "项目 DRI 书面接收"],
        ["T03 项目章程与 D0 确认", "G2 / 项目 DRI", "范围、里程碑、双方角色、双时钟、反馈 SLA、停表、验收", "客户确认 D0 日期"],
        ["T04 策略/范围基线", "G3 / 内容或实施负责人", "目标、受众/流程、范围内外、数量、质量、验收、开放问题", "验收人签署"],
        ["T05 内容卡与实验台账", "代运营周循环 / 内容负责人", "母内容 ID、80/20、假设、受众、CTA、事实来源、平台、结论", "每条可追溯"],
        ["T06 内部 QA 单", "G5 / QA", "评分、否决项、缺陷、返修、复验、批准版本", "满足业务线放行阈值"],
        ["T07 客户反馈与交付确认", "G6 / 项目 DRI", "版本、发送日、截止日、问题分类、轮次、验收、遗留项", "验收或视为验收证据"],
        ["T08 询盘与销售反馈表", "运营期 / 运营+客户销售", "来源内容、平台、时间、需求、LEAD0–4、跟进人、结果、无效原因", "周复盘可归因"],
        ["T09 周计划/周报", "G7 / 项目 DRI", "计划完成、质量、库存、数据、询盘、异常、等待、下周动作", "例会形成决策"],
        ["T10 异常单", "P0–P2 / 事件负责人；P3 进普通待办", "等级、时间线、影响、止损、根因、方案、更新、恢复验证、行动项", "批准关闭"],
        ["T11 变更申请单", "范围变化 / 项目 DRI", "新增/删除、原因、范围/时间/成本/质量影响、审批、实施版本", "批准后进入计划"],
        ["T12 月报/90 日结项", "D30/D90 / 项目 DRI", "交付、质量、双时钟、市场/询盘、成本、资产、续约/终止建议", "项目 DRI 签署；红线/补偿/终止时上收"],
    ], [1850, 1900, 3880, 1730], font_size=7.75, first_col_bold=True)

    add_heading(doc, "9.2 六张一页表的填写骨架", 1)
    add_table(doc, ["表单", "页首必填", "正文结构", "页尾签署"], [
        ["D0 确认单", "项目编号、客户、业务线/模式、合同版本、D0 日期", "启动条件逐项状态；缺口与替代；首个里程碑；双方 SLA", "项目 DRI、销售、客户接口人"],
        ["内容实验台账", "母内容 ID、平台、稳定/实验、观察窗", "目标受众；假设；单一变量；内容/CTA；发布证据；数据；结论", "内容负责人、QA"],
        ["询盘反馈表", "线索 ID、时间、平台、来源内容", "原始问题；LEAD0–4；客户/需求；销售接收时间；后续状态；无效原因", "运营、客户销售"],
        ["周报", "本周周期、负责人、项目灯号", "计划/完成；质量；库存；市场信号；询盘；等待；风险；下周 3 项动作", "项目 DRI、客户接口人"],
        ["异常单", "事件编号、等级、发现时间、报告人", "事实与影响；时间线；止损；原因；方案；恢复验证；防复发", "事件负责人、QA、总负责人"],
        ["变更单", "变更编号、提出方、日期、关联基线", "变更描述；必要性；范围/工期/费用/资源/风险影响；不变项", "客户批准人、商务、项目 DRI"],
    ], [1750, 2460, 3550, 1600], font_size=8.0, first_col_bold=True)

    add_heading(doc, "9.3 固定会议节奏", 1)
    add_table(doc, ["会议", "频率 / 时长", "只讨论", "会后 1 个工作日内"], [
        ["交付晨检", "工作日 10–15 分钟", "逾期、今日放行、输入缺口、红黄灯", "看板更新责任人与恢复点"],
        ["客户周会", "每周 30–45 分钟", "上周证据、策略/产品决策、客户待办、下周计划", "会议纪要和双方行动项"],
        ["质量校准", "每周 30 分钟", "典型通过/不通过样本、返修原因、模板规则", "QA 规则和示例库更新"],
        ["月度经营复盘", "每 30 日 60 分钟", "交付、质量、工时成本、经营信号、客户配合、续约风险", "继续/纠偏/扩围/停止决策"],
        ["重大异常会", "P0 即时；P1 当日", "事实、影响、止损、选项、外部口径、下次更新", "异常单和批准决策"],
    ], [1700, 1700, 4090, 1870], font_size=8.35, first_col_bold=True)

    add_heading(doc, "9.4 V1.0 上线计划", 1)
    add_table(doc, ["时间", "动作", "负责人", "验收"], [
        ["发布当日", "确认总负责人和各业务线 DRI；冻结 V1.0；建立统一项目主档目录", "总负责人", "制度发布记录；权限可用"],
        ["第 1–3 个工作日", "按第 7 章建项目主表和关联子表；配置红黄灯、逾期、80% 成本提醒", "项目 DRI / 系统管理员", "使用测试项目跑通提醒"],
        ["第 4–5 个工作日", "对销售、交付、内容、QA、技术进行 90 分钟培训；使用真实案例演练 G0–G6", "总负责人 / QA", "参训记录；演练问题闭环"],
        ["第 1–4 周", "选择 3 个项目试运行：优先同一细分行业，至少覆盖代运营，并尽量覆盖 SaaS/定制", "项目 DRI", "每周审计模板完整率与阶段门"],
        ["第 30 日", "复盘交付周期、一次通过率、客户等待、返工、成本、询盘回传和制度漏洞", "总负责人", "形成 V1.1 变更清单"],
        ["第 31 日起", "将通过验证的流程复制到全部新项目；老项目在下个里程碑切入", "各项目 DRI", "新项目 SOP 使用率 100%"],
    ], [1500, 4630, 1500, 1730], font_size=8.25, first_col_bold=True)

    add_heading(doc, "9.5 规模化健康指标", 1)
    add_table(doc, ["指标", "V1.0 建议目标", "用途"], [
        ["新项目 G0/G1/G2 使用率", "100%", "防止销售承诺与启动失控"],
        ["关键阶段证据完整率", "阶段门证据 100%；非关键审计字段 ≥95%", "确保可追溯、可验收"],
        ["里程碑按期率", "≥90%，外部等待单列", "衡量内部交付稳定性"],
        ["内容内部一次通过率", "通过条数÷首检条数；首月建基线，第 2 月目标由总负责人批准", "识别模板和培训问题"],
        ["严重异常复发率", "90 日同根因 P0/P1 数÷已关闭 P0/P1 数；目标逐季下降", "检验复盘是否有效"],
        ["客户销售反馈回传率", "试点期 ≥80%", "使询盘质量和续约判断有证据"],
        ["项目毛利/成本偏差", "所有项目可见；80% 预算预警执行率 100%", "防止规模扩大但亏损扩大"],
        ["交付周期 P50/P90", "按业务线、复杂度统计 D0 至上线/首发的中位数与 90 分位", "识别平均数掩盖的长尾"],
        ["阻塞老化", ">5 工作日项目数与占比每周可见；>10 日必须有重基线", "防止外部等待无限冻结"],
        ["组合容量利用率", "按角色每周统计；>85% 不得无批准新增启动", "保护 QA、技术和 DRI 瓶颈"],
    ], [2450, 2860, 4050], font_size=8.55, first_col_bold=True)

    add_heading(doc, "9.6 跨项目容量与 WIP 控制", 1)
    add_table(doc, ["控制项", "V1.0 规则", "红黄灯 / 决策"], [
        ["容量点", "S=1、M=2、L=4；X 不进入标准池，单独评估。30 日后按真实工时校准", "报价与 G0 必须同时填写未来 4 周需求点数"],
        ["可用容量", "按 DRI、内容、QA、技术分别统计可用工时并换算点数；休假/培训/维护先扣除", "任何关键角色无容量时不得只看总人数承诺"],
        ["利用率", "已承诺容量点÷未来 4 周可用容量点", "≤70% 绿；71–85% 黄并限制插单；>85% 红，原则上不接新启动"],
        ["WIP", "每人同时进行中的项目数和批次数由业务线负责人设上限；等待项目不占执行 WIP，但占恢复队列", "超过上限必须拆分、增援、延期或拒单"],
        ["优先级", "安全合规恢复 > 合同硬日期 > 已上线客户 > 已回款标准项目 > 试点/实验", "任何插单必须记录被挤压项目及新日期"],
        ["组合例会", "每周查看未来 4/8 周需求、瓶颈、阻塞老化、P50/P90 周期和毛利偏差", "由授权交付负责人做增援、重排、延期或拒单决策"],
    ], [1900, 4780, 2680], font_size=8.25, first_col_bold=True)

    add_heading(doc, "9.7 授权与上收矩阵", 1)
    add_table(doc, ["事项", "可下放批准", "必须上收总负责人"], [
        ["G0 准入", "S 项目由授权交付负责人；M 项目在标准价格/周期/边界内可批", "L/X、无容量、非标结果承诺、毛利低于底线"],
        ["范围变更", "不增价、不跨模块且不影响合同日的轻微变更，由项目 DRI 批", "增价/退款、跨模块、工期 >2 工作日、重大质量或合规影响"],
        ["QA 放行", "QA 对满足阈值的标准批次独立放行", "风险接受、P0/P1 恢复、重大事实/舆情影响"],
        ["异常", "P2/P3 由项目 DRI；P1 由授权业务线负责人先止损", "所有 P0、未控制 P1、数据/账号权属/违法违规"],
        ["补偿与终止", "项目 DRI 仅能提出方案", "退款、赔偿、免费扩围、长期暂停或合同终止"],
        ["代理", "每个 A 预先登记 1 名书面代理；代理期间权限等同但必须留痕", "无人有权时自动上收，不得多人共同签字代替唯一 A"],
    ], [1900, 4150, 3310], font_size=8.35, first_col_bold=True)
    add_callout(doc, "V1.0 使用要求", "本 SOP 先作为内部最低控制线。具体客户数量、平台规则、价格、SLA 和法律义务始终以最新合同与书面范围为准；任何不确定能力先 PoC，再进入标准销售口径。", "green")

    doc.core_properties.title = "灵枢 AI 内部商业化交付 SOP"
    doc.core_properties.subject = "标准 SaaS、企业定制/私有化与社媒代运营的内部交付控制体系"
    doc.core_properties.author = "灵枢 AI"
    doc.core_properties.keywords = "交付SOP, 社媒代运营, SaaS, 私有化, 80/20, 询盘, 阶段门"
    reorder_document_sections(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
